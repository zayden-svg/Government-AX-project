# main.py
import asyncio
import hashlib
import re
from datetime import datetime, date

import pandas as pd
from sqlalchemy import text
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from collectors import run_all_collectors
from biz_classifier import classify_and_score as biz_classify_and_score
from ai_utils import is_gemini_ready, analyze_postings_batch, AI_MAX_CONCURRENCY, AI_RPM_LIMIT
from db2 import get_engine, is_postgres

EXCEL_OUTPUT = "Gov-Tracker_결과.xlsx"

# ------------------------------------------------------------
# 공고 유형(입찰/RND/기타) - 이건 "AI 구분"과 다른 개념.
# 게시글 형식이 입찰공고문인지 R&D공모문인지만 가볍게 나누는 용도.
# ------------------------------------------------------------
RND_TYPE_KEYWORDS = ["과제공고", "지원사업", "R&D", "연구개발", "공모전", "지원과제", "사업 공모", "공모", "수요조사"]
BID_TYPE_KEYWORDS = ["입찰", "전자입찰", "구매", "용역", "발주", "제안서", "제안", "적격심사", "사전규격"]

TRACK_LABELS = {"RND": "R&D", "BIZ": "사업부", "": "미분류"}

# AI 분석이 실패했을 때만 쓰는 최후의 안전장치 (평소엔 사용 안 함)
AGENCY_TRACK_FALLBACK = {
    "IRIS": "RND", "KERIS": "RND", "NTIS": "RND", "TIPA": "RND",
    "KIAT": "RND", "INNOPOLIS": "RND", "국가AI전략위원회": "RND", "AIHub": "RND",
    "NIPA": "BIZ", "조달청": "BIZ", "행정안전부": "BIZ",
}


def classify_post_type(title: str) -> str:
    if not title:
        return "ETC"
    for kw in BID_TYPE_KEYWORDS:
        if kw in title:
            return "BID"
    for kw in RND_TYPE_KEYWORDS:
        if kw in title:
            return "RND"
    return "ETC"


# ------------------------------------------------------------
# 담당자명 정제
# ------------------------------------------------------------
AGENCY_SUFFIXES = [
    "센터", "재단", "진흥원", "위원회", "연구원", "정보원",
    "공사", "청", "부", "원", "팀", "실", "국", "협회", "본부",
]


def clean_manager_name(raw) -> str:
    if raw is None:
        return ""
    name = str(raw).strip()
    if not name:
        return ""
    if re.search(r"\d", name):
        return ""
    for suffix in AGENCY_SUFFIXES:
        if name.endswith(suffix):
            return ""
    if re.fullmatch(r"[가-힣]{2,4}", name):
        return name
    if re.fullmatch(r"[A-Za-z]{2,20}(\s[A-Za-z]{2,20})?", name):
        return name
    return ""


# ------------------------------------------------------------
# 중복 판별용 해시 (기관 + 제목 + 등록일 + 마감일 기준)
# ------------------------------------------------------------
def normalize_for_dedup(text_val) -> str:
    if not text_val:
        return ""
    t = str(text_val)
    t = re.sub(r"\s+", "", t)
    t = re.sub(r"[^\w가-힣]", "", t)
    return t.lower()


def build_dedup_hash(title: str, agency: str, reg_date: str, due_date: str) -> str:
    norm_title = normalize_for_dedup(title)
    norm_agency = normalize_for_dedup(agency)
    base = f"{norm_agency}|{norm_title}|{reg_date or ''}|{due_date or ''}"
    return hashlib.md5(base.encode("utf-8")).hexdigest()


def make_uniq_key(rec: dict) -> str:
    base = rec.get("url") or f"{rec.get('agency')}|{rec.get('title')}|{rec.get('reg_date')}"
    return hashlib.md5(base.encode("utf-8")).hexdigest()


# ------------------------------------------------------------
# DB 초기화 / 마이그레이션 (Supabase-Postgres, 로컬-SQLite 겸용)
# ------------------------------------------------------------
def init_db():
    engine = get_engine()
    with engine.begin() as conn:
        conn.execute(text("""
            CREATE TABLE IF NOT EXISTS postings (
                uniq_key TEXT PRIMARY KEY,
                source TEXT, agency TEXT, gubun TEXT, post_type TEXT, title TEXT,
                dept TEXT, manager TEXT, reg_date TEXT, due_date TEXT, budget TEXT,
                attach TEXT, views TEXT, url TEXT, grade TEXT, category TEXT,
                matched_keywords TEXT, recommended_solution TEXT, status TEXT,
                track TEXT, track_reason TEXT,
                ai_priority_score INTEGER, ai_priority_reason TEXT,
                dedup_hash TEXT,
                created_at TEXT, updated_at TEXT
            )
        """))
    ensure_columns(engine)
    with engine.begin() as conn:
        conn.execute(text("CREATE INDEX IF NOT EXISTS idx_dedup_hash ON postings(dedup_hash)"))
    return engine


def ensure_columns(engine):
    """예전 버전 DB에 새 컬럼(track_reason, ai_priority_score 등)이 없으면 추가"""
    needed = {
        "track": "TEXT",
        "track_reason": "TEXT",
        "ai_priority_score": "INTEGER",
        "ai_priority_reason": "TEXT",
        "dedup_hash": "TEXT",
    }
    with engine.begin() as conn:
        if is_postgres():
            existing = {row[0] for row in conn.execute(text(
                "SELECT column_name FROM information_schema.columns WHERE table_name='postings'"
            ))}
        else:
            existing = {row[1] for row in conn.execute(text("PRAGMA table_info(postings)"))}
        for col, coltype in needed.items():
            if col not in existing:
                conn.execute(text(f"ALTER TABLE postings ADD COLUMN {col} {coltype}"))


def find_existing_by_dedup_hash(engine, dedup_hash: str):
    """동일 dedup_hash를 가진 기존 레코드 조회 (AI 분석 결과 재사용 및 중복 판단에 사용)"""
    if not dedup_hash:
        return None
    with engine.begin() as conn:
        row = conn.execute(text("""
            SELECT uniq_key, agency, track, track_reason, ai_priority_score, ai_priority_reason
            FROM postings WHERE dedup_hash=:h LIMIT 1
        """), {"h": dedup_hash}).fetchone()
    return row


def upsert_posting(engine, rec: dict) -> str:
    uniq_key = rec.get("_uniq_key") or make_uniq_key(rec)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    params = {
        "uniq_key": uniq_key,
        "source": rec.get("source", ""), "agency": rec.get("agency", ""),
        "gubun": rec.get("gubun", ""), "post_type": rec.get("post_type", ""),
        "title": rec.get("title", ""), "dept": rec.get("dept", ""),
        "manager": rec.get("manager", ""), "reg_date": rec.get("reg_date", ""),
        "due_date": rec.get("due_date", ""), "budget": rec.get("budget", ""),
        "attach": rec.get("attach", ""), "views": rec.get("views", ""),
        "url": rec.get("url", ""), "grade": rec.get("grade", ""),
        "category": rec.get("category", ""), "matched_keywords": rec.get("matched_keywords", ""),
        "recommended_solution": rec.get("recommended_solution", ""),
        "track": rec.get("track", ""), "track_reason": rec.get("track_reason", ""),
        "ai_priority_score": rec.get("ai_priority_score"),
        "ai_priority_reason": rec.get("ai_priority_reason", ""),
        "dedup_hash": rec.get("dedup_hash", ""),
        "now": now,
    }

    with engine.begin() as conn:
        exists = conn.execute(text("SELECT 1 FROM postings WHERE uniq_key=:uniq_key"), params).fetchone()

        if exists:
            conn.execute(text("""
                UPDATE postings SET
                    source=:source, agency=:agency, gubun=:gubun, post_type=:post_type, title=:title,
                    dept=:dept, manager=:manager, reg_date=:reg_date, due_date=:due_date, budget=:budget,
                    attach=:attach, views=:views, url=:url, grade=:grade, category=:category,
                    matched_keywords=:matched_keywords, recommended_solution=:recommended_solution,
                    track=:track, track_reason=:track_reason,
                    ai_priority_score=:ai_priority_score, ai_priority_reason=:ai_priority_reason,
                    dedup_hash=:dedup_hash, status='진행중', updated_at=:now
                WHERE uniq_key=:uniq_key
            """), params)
            return "updated"
        else:
            conn.execute(text("""
                INSERT INTO postings (
                    uniq_key, source, agency, gubun, post_type, title, dept, manager,
                    reg_date, due_date, budget, attach, views, url,
                    grade, category, matched_keywords, recommended_solution,
                    track, track_reason, ai_priority_score, ai_priority_reason, dedup_hash,
                    status, created_at, updated_at
                ) VALUES (
                    :uniq_key, :source, :agency, :gubun, :post_type, :title, :dept, :manager,
                    :reg_date, :due_date, :budget, :attach, :views, :url,
                    :grade, :category, :matched_keywords, :recommended_solution,
                    :track, :track_reason, :ai_priority_score, :ai_priority_reason, :dedup_hash,
                    '진행중', :now, :now
                )
            """), params)
            return "new"


def mark_expired(engine):
    today = date.today().strftime("%Y-%m-%d")
    with engine.begin() as conn:
        conn.execute(text("""
            UPDATE postings SET status='마감'
            WHERE due_date != '' AND due_date < :today AND status != '마감'
        """), {"today": today})


def _apply_grade_style(path):
    wb = load_workbook(path)
    ws = wb["전체"]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    orange_fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")

    headers = [cell.value for cell in ws[1]]
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    if "등급" in headers:
        grade_col = headers.index("등급") + 1
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            grade_cell = row[grade_col - 1]
            if grade_cell.value == "상":
                for c in row:
                    c.fill = red_fill
            elif grade_cell.value == "중":
                for c in row:
                    c.fill = orange_fill

    if "제목" in headers and "원문링크" in headers:
        title_col = headers.index("제목") + 1
        url_col = headers.index("원문링크") + 1
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            title_cell = row[title_col - 1]
            url_cell = row[url_col - 1]
            if url_cell.value:
                title_cell.hyperlink = url_cell.value
                title_cell.font = Font(color="0563C1", underline="single")
        ws.delete_cols(url_col)

    for col_cells in ws.columns:
        max_len = max((len(str(c.value)) if c.value else 0) for c in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)

    ws.freeze_panes = "A2"
    wb.save(path)


def export_to_excel(engine):
    df = pd.read_sql_query("""
        SELECT
            agency AS 기관, track AS 분류, track_reason AS AI구분근거,
            post_type AS 공고유형, gubun AS 세부구분, title AS 제목,
            dept AS 담당부서, manager AS 담당자, reg_date AS 등록일, due_date AS 마감일,
            grade AS 등급, category AS 카테고리, recommended_solution AS 추천솔루션,
            matched_keywords AS 매칭키워드, ai_priority_score AS AI연관도점수,
            ai_priority_reason AS AI연관도근거, status AS 상태, url AS 원문링크
        FROM postings
        ORDER BY reg_date DESC
    """, engine)

    df = df.fillna("")
    df["분류"] = df["분류"].map(lambda v: TRACK_LABELS.get(v, v) if v else "미분류")

    with pd.ExcelWriter(EXCEL_OUTPUT, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="전체", index=False)

    _apply_grade_style(EXCEL_OUTPUT)
    print(f"[완료] 결과 파일 저장: {EXCEL_OUTPUT} ({len(df)}건)")


def _build_ai_info_block(rec: dict, agency: str) -> str:
    return (
        f"제목: {rec.get('title', '')}\n"
        f"기관: {rec.get('agency', agency)}\n"
        f"부서: {rec.get('dept', '')}\n"
        f"공고유형(원본 구분값): {rec.get('gubun', '')}\n"
        f"본문 일부: {str(rec.get('content', ''))[:1000]}"
    )


def collect_and_process():
    engine = init_db()
    new_count = updated_count = skipped_count = duplicate_count = ai_called_count = 0

    results_by_agency = run_all_collectors(limit=10)
    total_records = sum(len(v) for v in results_by_agency.values())
    processed = 0

    print(f"\n[1단계] 총 {total_records}건 수집 완료. 등급판정/중복확인 진행 중...\n")

    ready_records = []           # 바로 DB 저장 가능한 레코드 (재사용 or 이번에 AI 불필요)
    pending_ai = []               # [(uniq_key, info_block), ...] - 새로 AI 분석이 필요한 건
    rec_map = {}                  # uniq_key -> rec (AI 결과를 나중에 채워 넣기 위함)

    for agency, records in results_by_agency.items():
        for rec in records:
            processed += 1
            if not rec.get("title"):
                skipped_count += 1
                continue

            rec["manager"] = clean_manager_name(rec.get("manager", ""))
            rec["post_type"] = classify_post_type(rec["title"])

            score_info = biz_classify_and_score(rec["title"], rec.get("content", ""))
            rec["grade"] = score_info["등급"]
            rec["category"] = score_info["카테고리"]
            rec["recommended_solution"] = score_info["추천솔루션"]
            rec["matched_keywords"] = score_info["매칭키워드"]

            dedup_hash = build_dedup_hash(
                rec.get("title", ""), rec.get("agency", agency),
                rec.get("reg_date", ""), rec.get("due_date", "")
            )
            rec["dedup_hash"] = dedup_hash
            current_uniq_key = make_uniq_key(rec)
            rec["_uniq_key"] = current_uniq_key

            existing = find_existing_by_dedup_hash(engine, dedup_hash)

            if existing and existing[0] != current_uniq_key:
                duplicate_count += 1
                print(f"  [{processed}/{total_records}] 중복 제외: {rec['title'][:40]}")
                continue

            # 기준을 track(문자열) 대신 ai_priority_score(숫자)로 바꿈:
            # AI 분석이 "실패"했던 건은 score가 None이라서 여기 안 걸리고 다시 AI 분석 대상으로 감.
            if existing and existing[0] == current_uniq_key and existing[4] is not None:
                rec["track"] = existing[2]
                rec["track_reason"] = existing[3]
                rec["ai_priority_score"] = existing[4]
                rec["ai_priority_reason"] = existing[5]
                print(f"  [{processed}/{total_records}] 기존 AI 분석 재사용: {rec['title'][:40]}")
                ready_records.append(rec)
            else:
                pending_ai.append((current_uniq_key, _build_ai_info_block(rec, agency)))
                rec_map[current_uniq_key] = rec

    ai_total = len(pending_ai)
    print(
        f"\n[2단계] 신규 AI 분석 대상 {ai_total}건. "
        f"동시 {AI_MAX_CONCURRENCY}건씩, 분당 최대 {AI_RPM_LIMIT}건 속도로 병렬 처리합니다...\n"
    )

    ai_results = {}
    if ai_total > 0 and is_gemini_ready():
        done_counter = {"n": 0}

        def _progress(key, result):
            done_counter["n"] += 1
            title = rec_map[key].get("title", "")[:40]
            if result and not result.get("error"):
                print(f"  [{done_counter['n']}/{ai_total}] AI 분석 완료: {result['track']} / {result['score']}점 - {title}")
            else:
                err = result.get("error", "알수없음") if result else "알수없음"
                print(f"  [{done_counter['n']}/{ai_total}] AI 분석 실패: {err} - {title}")

        ai_results = asyncio.run(analyze_postings_batch(pending_ai, progress_cb=_progress))
        ai_called_count = ai_total
    elif ai_total > 0:
        print("  [경고] Gemini API 미설정으로 AI 분석을 건너뜁니다. 기관 기본값으로 임시 분류됩니다.")

    for uniq_key, rec in rec_map.items():
        ai_result = ai_results.get(uniq_key)
        if ai_result and not ai_result.get("error"):
            rec["track"] = ai_result["track"]
            rec["track_reason"] = ai_result["track_reason"]
            rec["ai_priority_score"] = ai_result["score"]
            rec["ai_priority_reason"] = ai_result["score_reason"]
        else:
            err_msg = ai_result.get("error", "Gemini 미설정") if ai_result else "Gemini 미설정"
            rec["track"] = AGENCY_TRACK_FALLBACK.get(rec.get("agency", ""), "BIZ")
            rec["track_reason"] = f"AI 분석 실패({err_msg})로 기관 기본값으로 임시 분류됨"
            rec["ai_priority_score"] = None
            rec["ai_priority_reason"] = "AI 분석 대기 중"
        ready_records.append(rec)

    print(f"\n[3단계] DB 저장 중... (총 {len(ready_records)}건)\n")
    for rec in ready_records:
        status = upsert_posting(engine, rec)
        if status == "new":
            new_count += 1
        elif status == "updated":
            updated_count += 1

    mark_expired(engine)
    export_to_excel(engine)

    print(
        f"\n신규 {new_count}건 / 갱신 {updated_count}건 / "
        f"제목없음 제외 {skipped_count}건 / 중복 제외 {duplicate_count}건 / "
        f"AI 신규 분석 {ai_called_count}건 처리 완료"
    )


if __name__ == "__main__":
    collect_and_process()
