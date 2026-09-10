# keyword_config_생성.py (최초 1회 실행용)
import openpyxl
from openpyxl.styles import Font, PatternFill

wb = openpyxl.Workbook()

# ── 시트1: 공고유형 (R&D 과제 vs 영업 입찰공고) ──
ws1 = wb.active
ws1.title = "공고유형"
ws1.append(["유형", "키워드", "설명"])
type_rows = [
    ("RND", "과제공고", "R&D 지원사업 공모"),
    ("RND", "지원사업", ""),
    ("RND", "공모", ""),
    ("RND", "R&D", ""),
    ("RND", "연구개발", ""),
    ("RND", "출연금", ""),
    ("RND", "기술개발사업", ""),
    ("RND", "실증사업", ""),
    ("RND", "지원과제", ""),
    ("BID", "입찰공고", "구매/용역 입찰"),
    ("BID", "제안요청서", ""),
    ("BID", "구매", ""),
    ("BID", "용역", ""),
    ("BID", "구축", ""),
    ("BID", "전자입찰", ""),
    ("BID", "협상에의한계약", ""),
    ("BID", "사전규격", ""),
]
for row in type_rows:
    ws1.append(row)

# ── 시트2: 제품분류 (기존 biz_classifier 내용 이전) ──
ws2 = wb.create_sheet("제품분류")
ws2.append(["카테고리", "우선순위", "키워드", "추천솔루션"])
cat_rows = [
    ("업무유형(예약·청약·접수·시험)", "1차", "예매", "NetFUNNEL"),
    ("업무유형(예약·청약·접수·시험)", "1차", "발권", "NetFUNNEL"),
    ("업무유형(예약·청약·접수·시험)", "1차", "수강신청", "NetFUNNEL"),
    ("업무유형(예약·청약·접수·시험)", "1차", "접수기간", "NetFUNNEL"),
    ("업무유형(예약·청약·접수·시험)", "1차", "채용시험 접수", "NetFUNNEL"),
    ("트래픽·접속제어", "1차", "대기열", "NetFUNNEL"),
    ("트래픽·접속제어", "1차", "대기시스템", "NetFUNNEL"),
    ("트래픽·접속제어", "1차", "동시접속", "NetFUNNEL"),
    ("매크로·부정접속방어", "1차", "매크로", "MBUSTER"),
    ("매크로·부정접속방어", "1차", "부정접속", "MBUSTER"),
    ("매크로·부정접속방어", "1차", "어뷰징", "MBUSTER"),
    ("시스템구축·전환", "1차", "차세대", "NetFUNNEL"),
    ("시스템구축·전환", "1차", "고도화", "NetFUNNEL"),
    ("시스템구축·전환", "1차", "통합플랫폼", "NetFUNNEL"),
    ("보조신호(결합시유효)", "2차", "홈페이지 구축", "NetFUNNEL/MBUSTER 공통 검토"),
    ("보조신호(결합시유효)", "2차", "인프라 확충", "NetFUNNEL/MBUSTER 공통 검토"),
    ("보조신호(결합시유효)", "2차", "클라우드 전환", "NetFUNNEL/MBUSTER 공통 검토"),
    ("보조신호(결합시유효)", "2차", "안정화", "NetFUNNEL/MBUSTER 공통 검토"),
    ("보조신호(결합시유효)", "2차", "DR", "NetFUNNEL/MBUSTER 공통 검토"),
]
for row in cat_rows:
    ws2.append(row)

for ws in (ws1, ws2):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4E78", fill_type="solid")

wb.save("keyword_config.xlsx")
print("keyword_config.xlsx 생성 완료")
