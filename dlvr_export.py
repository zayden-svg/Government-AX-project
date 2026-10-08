# dlvr_export.py
# 조달청 쇼핑몰(디지털서비스몰·종합쇼핑몰) 납품요구 수집 → "어느 기관이 언제 무엇을 샀나" 엑셀
#   · 자사(NetFUNNEL·NFA·MBUSTER)·경쟁사(xQueue·DynaPath·에버세이프 등)·비교군 제품의 기관 구매 이력
#   · API: 조달청_나라장터쇼핑몰 품목정보 서비스 / getDlvrReqDtlInfoList (납품요구 상세)
#   · 공공데이터포털 활용신청 필요(자동승인). 하루 1,000회 한도 — 이 수집은 수십~수백 회면 끝남
#
# 실행: python dlvr_export.py --out dlvr.xlsx
import argparse
import io
import json
import os
import re
import sys
import time
from datetime import datetime, timedelta
from urllib.parse import unquote

import pandas as pd
import requests

BASE = "https://apis.data.go.kr/1230000/at/ShoppingMallPrdctInfoService"
OP_DTL = f"{BASE}/getDlvrReqDtlInfoList"
ROWS = 999

# 디지털서비스몰 계약상품정보(2026-03 추출본)에서 찾은 계약번호 — (계약번호, 구분, 제품군, 제조사, 제품)
TARGET_CONTRACTS = [
    ("002150153", "자사", "NF", "에스티씨랩", "NetFUNNEL v3.0"),
    ("002460509", "자사", "NF", "에스티씨랩", "NetFUNNEL v3.0"),
    ("R26TA01343941", "자사", "NF", "에스티씨랩", "NetFUNNEL v4.0"),
    ("R26TA01375633", "자사", "NFA", "에스티씨랩(에티버스이피에이 판매)", "API NetFUNNEL v4.0"),
    ("002460506", "자사", "BM/MB", "에스티씨랩", "MBUSTER v2.0"),
    ("002461613", "경쟁사", "대기열·유량제어", "소프트베이스", "xQueue v1.0"),
    ("002461139", "경쟁사", "봇·매크로 차단", "스크립터스", "DynaPath 3.1"),
    ("R25TA00251499", "경쟁사", "봇·매크로 차단", "에버스핀", "에버세이프 웹 보안솔루션 V4.2.1"),
    ("R25TA00145277", "경쟁사", "앱 위변조 방지", "에버스핀", "에버세이프 v2.0"),
    ("002161010", "경쟁사", "앱 위변조 방지", "에버스핀(가온아이 판매)", "에버세이프 v2.0"),
    ("R25TA00466390", "비교군", "봇·매크로 차단", "네스토리", "BotfenderAI V2.0"),
    ("002162591", "비교군", "API 게이트웨이", "메타빌드", "MESIM APIG v1.5"),
    ("002462107", "비교군", "API 게이트웨이", "메타빌드", "MESIM APIG v1.5"),
    ("002261452", "비교군", "API 게이트웨이", "이데아텍", "i-ONE API Gateway v1.2"),
    ("R25TA01041570", "비교군", "API 게이트웨이", "메가투스", "Megaapim V1.0"),
]
# 계약번호로 못 잡는 예전·다른 계약 보완: 물품규격명(제품명) 월별 검색 — (검색어, 구분, 제품군, 시작연도)
#   조달청 API는 날짜 조회를 한 달 단위로만 허용 → 한 달씩 훑음
SCAN_NAMES = [
    ("에스티씨랩", "자사", "자사 제품", 2012), ("NetFUNNEL", "자사", "NF", 2012),
    ("xQueue", "경쟁사", "대기열·유량제어", 2020), ("DynaPath", "경쟁사", "봇·매크로 차단", 2020),
    ("에버세이프", "경쟁사", "봇·앱 보안", 2020),
]
RENEWAL_BEFORE = "2020-01-01"   # 자사 고객 중 마지막 구매가 이 날짜 이전이고 그 뒤 구매가 없는 기관 = 리뉴얼 타겟
MY_REGIONS = ["서울", "인천", "강원", "전북", "전남", "광주", "제주"]   # 담당 지역 표시용


def _key():
    k = (os.getenv("G2B_SERVICE_KEY") or "").strip()
    return unquote(k) if "%" in k else k


def mask(s):
    return re.sub(r"(serviceKey|key)=[^&\s'\"]+", r"\1=***", str(s))


class ApiError(Exception):
    pass


CALLS = {"n": 0}


def call(params):
    q = {"serviceKey": _key(), "type": "json", "numOfRows": str(ROWS), **params}
    last = "429 반복"
    for attempt in range(4):
        CALLS["n"] += 1
        try:
            r = requests.get(OP_DTL, params=q, headers={"User-Agent": "Mozilla/5.0"}, timeout=(15, 60))
        except requests.RequestException as e:
            time.sleep(5 * (attempt + 1))
            last = e
            continue
        if r.status_code == 429:
            time.sleep(30 * (attempt + 1))
            continue
        txt = r.text or ""
        if "SERVICE_KEY_IS_NOT_REGISTERED" in txt:
            raise ApiError("활용신청이 아직 승인되지 않음 (조달청_나라장터쇼핑몰 품목정보 서비스)")
        if "LIMITED_NUMBER" in txt:
            raise ApiError("하루 호출 한도 초과")
        try:
            data = r.json()
        except ValueError:
            raise ApiError(mask(txt[:200]))
        err = data.get("nkoneps.com.response.ResponseError")
        if err:
            h = err.get("header") or {}
            raise ApiError(f"요청값 오류 {h.get('resultCode')} {h.get('resultMsg')}")
        root = data.get("response") or {}
        head = root.get("header") or {}
        if str(head.get("resultCode", "00")) not in ("00", "0"):
            raise ApiError(f"{head.get('resultCode')} {head.get('resultMsg')}")
        body = root.get("body") or {}
        items = body.get("items") or []
        if isinstance(items, dict):
            items = items.get("item") or []
        if isinstance(items, dict):
            items = [items]
        return items, int(body.get("totalCount") or 0)
    raise ApiError(f"접속 실패: {mask(last)}")


def fetch_all(params):
    out, page = [], 1
    while True:
        items, total = call({**params, "pageNo": str(page)})
        out.extend(items)
        if not items or len(out) >= total or len(items) < ROWS:
            return out
        page += 1


def month_chunks(start_year):
    cur = datetime(start_year, 1, 1)
    end = datetime.now()
    while cur <= end:
        nxt = (cur.replace(day=28) + timedelta(days=4)).replace(day=1)
        yield cur.strftime("%Y%m%d"), min(nxt - timedelta(days=1), end).strftime("%Y%m%d")
        cur = nxt


def _fam(spec):
    t = str(spec)
    if "API NetFUNNEL" in t:
        return "NFA"
    if "NetFUNNEL" in t or "넷퍼넬" in t:
        return "NF"
    if "MBUSTER" in t or "엠버스터" in t or "봇매니저" in t:
        return "BM/MB"
    return ""


PARTIAL = {"stopped": ""}
OFFLINE = False                   # True면 API를 부르지 않고 캐시 파일만으로 결과 구성 (통합 엑셀용)
CACHE_FILE = "raw_cache.json"     # 지난 실행에서 받아 둔 결과 (지난달 이전 자료는 바뀌지 않으므로 다시 안 부름)
WORKERS = 6


def _load_cache():
    try:
        return json.load(open(CACHE_FILE, encoding="utf-8"))
    except Exception:
        return {}


def _run_tasks(tasks, cache, log_label):
    """tasks: [(cache_key, params)] → 캐시에 없는 것만 동시에 조회. 한도 초과 시 ApiError"""
    from concurrent.futures import ThreadPoolExecutor, as_completed
    if OFFLINE:
        return
    todo = [(k, p) for k, p in tasks if k not in cache]
    stop = {"err": None}

    def one(k, p):
        if stop["err"]:
            return k, None
        try:
            return k, fetch_all(p)
        except ApiError as e:
            stop["err"] = e
            return k, None
    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        for f in as_completed([ex.submit(one, k, p) for k, p in todo]):
            k, items = f.result()
            if items is not None:
                cache[k] = items
    print(f"[{log_label}] 새로 조회 {len(todo)}건 중 {sum(1 for k, _ in todo if k in cache)}건 완료 · 누적 API {CALLS['n']}회", flush=True)
    if stop["err"]:
        raise stop["err"]


def collect(years=None):
    rows, log = [], []
    cache = _load_cache()
    this_month = datetime.now().strftime("%Y%m")
    # 계약번호 조회와 이번 달 검색은 매번 새로 (새 구매 반영)
    if not OFFLINE:
        cache = {k: v for k, v in cache.items() if not (k.startswith("c|") or k.split("|")[-2][:6] == this_month)}
    try:
        _collect(rows, log, cache)
    except ApiError as e:            # 한도 초과 등 → 모은 데까지 엑셀로 만들고, 다음 실행 때 이어서
        PARTIAL["stopped"] = str(e)
        log.append(f"[중단] {e} — 여기까지 모은 결과로 엑셀 생성 (다음 실행 때 이어서)")
    if not OFFLINE:
        json.dump(cache, open(CACHE_FILE, "w", encoding="utf-8"), ensure_ascii=False)
    print("\n".join(log))
    return rows


def _rows_from(cache, key, kind, fam, maker, prod, how):
    return [{**it, "_구분": kind, "_제품군": _fam(it.get("prdctIdntNoNm")) or fam, "_제조사": maker,
             "_대상제품": prod, "_찾은방법": how} for it in cache.get(key, [])]


def _collect(rows, log, cache):
    # ① 알려진 계약번호 → 그 계약으로 들어온 납품요구 전부 (기간 제한 없음)
    tasks = [(f"c|{no}", {"inqryDiv": "3", "cntrctNo": no}) for no, *_ in TARGET_CONTRACTS]
    try:
        _run_tasks(tasks, cache, "계약번호")
    finally:
        for no, kind, fam, maker, prod in TARGET_CONTRACTS:
            got = _rows_from(cache, f"c|{no}", kind, fam, maker, prod, f"계약번호 {no}")
            rows += got
            log.append(f"계약 {no} {prod}: {len(got)}건")
    # ② 제품명으로 한 달씩 검색 (예전 계약·목록에 없는 계약 찾기)
    tasks, meta = [], {}
    for nm, kind, fam, y0 in SCAN_NAMES:
        for b, e in month_chunks(y0):
            k = f"m|{nm}|{b}|{e}"
            tasks.append((k, {"inqryDiv": "1", "inqryBgnDate": b, "inqryEndDate": e, "prdctIdntNoNm": nm}))
            meta[k] = (nm, kind, fam)
    try:
        _run_tasks(tasks, cache, "제품명 월별 검색")
    finally:
        found_new, cnt = {}, {}
        known = {no for no, *_ in TARGET_CONTRACTS}
        for k, (nm, kind, fam) in meta.items():
            got = _rows_from(cache, k, kind, fam, "", nm, f"제품명 '{nm}'")
            rows += got
            cnt[nm] = cnt.get(nm, 0) + len(got)
            for it in got:
                no = it.get("cntrctNo")
                if no and no not in known:
                    found_new[no] = (kind, fam)
        log += [f"제품명 {nm}: {n}건" for nm, n in cnt.items()]
    # ③ 새로 찾은 계약번호도 전부 조회
    tasks = [(f"c|{no}", {"inqryDiv": "3", "cntrctNo": no}) for no in found_new]
    try:
        _run_tasks(tasks, cache, "추가 계약번호")
    finally:
        for no, (kind, fam) in found_new.items():
            got = _rows_from(cache, f"c|{no}", kind, fam, "", "", f"계약번호 {no}(검색으로 발견)")
            rows += got
            log.append(f"추가 계약 {no}: {len(got)}건")


def _num(v):
    try:
        return int(float(str(v).replace(",", "")))
    except (TypeError, ValueError):
        return None


def _date(v):
    s = re.sub(r"[^0-9]", "", str(v or ""))[:8]
    return f"{s[:4]}-{s[4:6]}-{s[6:8]}" if len(s) == 8 else ""


def build(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    df = pd.DataFrame(rows).fillna("")
    if df.empty:
        raise SystemExit("[SKIP] 납품요구 결과 0건")
    # 같은 납품요구·물품은 최신 변경차수만
    df["_chg"] = pd.to_numeric(df.get("dlvrReqChgOrd", 0), errors="coerce").fillna(0)
    df["_src"] = df["_찾은방법"].str.startswith("계약번호").map({True: 0, False: 1})
    df = (df.sort_values(["_chg", "_src"], ascending=[False, True])
            .drop_duplicates(subset=["dlvrReqNo", "prdctSno", "prdctIdntNo"], keep="first"))
    # 대상 제품 이름은 실제 규격명에서 (예: 'NetFUNNEL v3.0')
    df["_대상제품"] = [(str(sp).split(",")[2].strip() if str(sp).count(",") >= 2 else t) for sp, t in zip(df["prdctIdntNoNm"], df["_대상제품"])]
    df["납품요구일"] = df["dlvrReqRcptDate"].map(_date)
    df["금액"] = df["prdctAmt"].map(_num)
    df["유지보수·재구매 예상"] = (pd.to_datetime(df["납품요구일"], errors="coerce") + pd.Timedelta(days=365)).dt.strftime("%Y-%m")
    df["담당지역"] = df["dminsttRgnNm"].map(lambda r: "Y" if any(k in str(r) for k in MY_REGIONS) else "")
    order = {"자사": 0, "경쟁사": 1, "비교군": 2}
    df = df.assign(_o=df["_구분"].map(order)).sort_values(["_o", "납품요구일"], ascending=[True, False])

    cols = [("구분", "_구분", 7), ("제품군", "_제품군", 14), ("납품요구일", "납품요구일", 11), ("수요기관", "dminsttNm", 30),
            ("기관구분", "dmndInsttDivNm", 10), ("지역", "dminsttRgnNm", 14), ("담당지역", "담당지역", 7),
            ("납품요구건명", "dlvrReqNm", 40), ("물품규격(제품)", "prdctIdntNoNm", 50), ("수량", "prdctQty", 6),
            ("금액(원)", "금액", 13), ("판매업체", "corpNm", 20), ("납품기한", "dlvrTmlmtDate", 11),
            ("유지보수·재구매 예상", "유지보수·재구매 예상", 12), ("납품요구번호", "dlvrReqNo", 14), ("계약번호", "cntrctNo", 14),
            ("찾은 방법", "_찾은방법", 18)]
    wb = Workbook()
    navy = PatternFill("solid", fgColor="1B2A4A")
    hf = Font(bold=True, color="FFFFFF")

    def sheet(ws, d, cdefs):
        ws.append([t for t, _, _ in cdefs])
        for i, (_, _, w) in enumerate(cdefs, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
            ws.cell(1, i).font, ws.cell(1, i).fill = hf, navy
        for r in d.to_dict("records"):
            ws.append([(_date(r.get(c)) if c == "dlvrTmlmtDate" else r.get(c, "")) for _, c, _ in cdefs])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cdefs))}{len(d) + 1}"
        amt_i = [t for t, _, _ in cdefs].index("금액(원)") + 1 if "금액(원)" in [t for t, _, _ in cdefs] else None
        if amt_i:
            for row in ws.iter_rows(min_row=2, min_col=amt_i, max_col=amt_i):
                for c in row:
                    c.number_format = "#,##0"

    # 기관별 요약
    g = df.groupby(["_구분", "dminsttNm"])
    summ = pd.DataFrame({
        "제품": g["_대상제품"].agg(lambda s: ", ".join(dict.fromkeys(s))),
        "구매 건수": g["dlvrReqNo"].nunique(),
        "금액 합계(원)": g["금액"].sum(min_count=1),
        "최근 구매일": g["납품요구일"].max(),
        "지역": g["dminsttRgnNm"].first(),
        "담당지역": g["담당지역"].first(),
        "판매업체": g["corpNm"].agg(lambda s: ", ".join(dict.fromkeys(s))),
    }).reset_index().rename(columns={"_구분": "구분", "dminsttNm": "수요기관"})
    summ["다음 접촉 시점"] = (pd.to_datetime(summ["최근 구매일"], errors="coerce") + pd.Timedelta(days=300)).dt.strftime("%Y-%m")
    summ = summ.assign(_o=summ["구분"].map(order)).sort_values(["_o", "최근 구매일"], ascending=[True, False]).drop(columns="_o")

    ws = wb.active
    ws.title = "안내"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 100
    info = [
        ("쇼핑몰 납품요구 현황 (자사·경쟁사 제품)", ""),
        ("기준일", datetime.now().strftime("%Y-%m-%d")),
        ("출처", "조달청_나라장터쇼핑몰 품목정보 서비스 · 납품요구 상세 (공공데이터포털 OpenAPI) — 자사 2012년~, 경쟁사 2020년~"),
        ("대상", "디지털서비스몰에 등록된 자사·경쟁사·비교군 제품의 계약번호 + 제품명 검색"),
        ("", ""),
        ("시트", "내용"),
        ("리뉴얼 타겟", "자사 제품을 2020년 이전에 산 뒤 그 이후 구매 이력이 없는 기관 (마지막 구매일 최신순)"),
        ("기관별 요약", "기관마다 산 제품·건수·금액·최근 구매일 → 다음 접촉 시점(최근 구매 + 10개월)"),
        ("경쟁사 고객", "경쟁사·비교군 제품을 산 기관 = 교체 영업 대상"),
        ("자사 고객", "자사 제품을 산 기관 = 유지보수·추가 구매 대상"),
        ("전체 내역", "납품요구 한 건 한 건 (변경차수는 최신만)"),
        ("", ""),
        ("참고", "유지보수·재구매 예상 = 구매일 + 1년 (추정). 담당지역 = 서울·인천·강원·전라·광주·제주"),
    ]
    for a, b in info:
        ws.append([a, b])
    ws["A1"].font = Font(bold=True, size=15)
    for c in ("A6", "B6"):
        ws[c].font, ws[c].fill = hf, navy

    scols = [("구분", "구분", 7), ("수요기관", "수요기관", 32), ("제품", "제품", 36), ("구매 건수", "구매 건수", 8),
             ("금액(원)", "금액 합계(원)", 14), ("최근 구매일", "최근 구매일", 11), ("다음 접촉 시점", "다음 접촉 시점", 12),
             ("지역", "지역", 14), ("담당지역", "담당지역", 7), ("판매업체", "판매업체", 30)]
    own = summ[summ["구분"] == "자사"]
    renew = own[(own["최근 구매일"] != "") & (own["최근 구매일"] < RENEWAL_BEFORE)].sort_values("최근 구매일", ascending=False)
    rcols = [("수요기관", "수요기관", 32), ("마지막 구매일", "최근 구매일", 12), ("산 제품", "제품", 36), ("구매 건수", "구매 건수", 8),
             ("금액(원)", "금액 합계(원)", 14), ("지역", "지역", 14), ("담당지역", "담당지역", 7), ("판매업체", "판매업체", 30)]
    sheet(wb.create_sheet("리뉴얼 타겟", 1), renew, rcols)
    sheet(wb.create_sheet("기관별 요약"), summ, scols)
    sheet(wb.create_sheet("경쟁사 고객"), summ[summ["구분"] != "자사"], scols)
    sheet(wb.create_sheet("자사 고객"), summ[summ["구분"] == "자사"], scols)
    sheet(wb.create_sheet("전체 내역"), df, cols)
    buf = io.BytesIO()
    wb.save(buf)
    meta = {"renewal": len(renew), "rows": len(df), "insttn": int(summ["수요기관"].nunique()),
            "own": int((summ["구분"] == "자사").sum()), "comp": int((summ["구분"] != "자사").sum())}
    return buf.getvalue(), meta


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="dlvr.xlsx")
    ap.add_argument("--years", type=int, default=0, help="(사용 안 함 — 제품별 시작연도는 SCAN_NAMES)")
    a = ap.parse_args()
    rows = collect()
    if not rows:
        print(f"[FAIL] 결과 없음 {PARTIAL['stopped']} · API {CALLS['n']}회")
        return 1
    data, meta = build(rows)
    meta["stopped"] = PARTIAL["stopped"]
    open(a.out, "wb").write(data)
    print(f"[OK] 납품요구 엑셀: API {CALLS['n']}회 · {json.dumps(meta, ensure_ascii=False)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
