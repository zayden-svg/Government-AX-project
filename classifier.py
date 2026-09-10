# classifier.py
import openpyxl
from pathlib import Path

KEYWORD_FILE = "keyword_config.xlsx"


def _load_type_keywords():
    wb = openpyxl.load_workbook(KEYWORD_FILE, data_only=True)
    ws = wb["공고유형"]
    rnd_kw, bid_kw = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        유형, 키워드 = row[0], row[1]
        if not 키워드:
            continue
        if 유형 == "RND":
            rnd_kw.append(str(키워드))
        elif 유형 == "BID":
            bid_kw.append(str(키워드))
    return rnd_kw, bid_kw


def _load_category_keywords():
    wb = openpyxl.load_workbook(KEYWORD_FILE, data_only=True)
    ws = wb["제품분류"]
    categories = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        cat, priority, kw, solution = row[0], row[1], row[2], row[3]
        if not cat or not kw:
            continue
        categories.setdefault(cat, {"1차": [], "2차": [], "solution": solution})
        if priority == "1차":
            categories[cat]["1차"].append(str(kw))
        else:
            categories[cat]["2차"].append(str(kw))
        if solution:
            categories[cat]["solution"] = solution
    return categories


def classify_post_type(title: str) -> str:
    """제목에 RND 키워드가 있으면 'RND', BID 키워드가 있으면 'BID', 둘 다 없으면 'BID'를 기본값으로."""
    rnd_kw, bid_kw = _load_type_keywords()
    text = title or ""
    if any(kw in text for kw in rnd_kw):
        return "RND"
    if any(kw in text for kw in bid_kw):
        return "BID"
    return "BID"  # 판단 불가 시 입찰공고 쪽으로 기본 분류 (영업팀이 놓치지 않도록)


def classify_and_score(title: str, content: str = "") -> dict:
    categories = _load_category_keywords()
    text = f"{title} {content}"
    matched_1st, matched_2nd, matched_cats, solutions = [], [], set(), set()

    for cat, info in categories.items():
        for kw in info["1차"]:
            if kw in text:
                matched_1st.append(kw); matched_cats.add(cat)
                if info["solution"]:
                    solutions.add(info["solution"])
        for kw in info["2차"]:
            if kw in text:
                matched_2nd.append(kw); matched_cats.add(cat)
                if info["solution"]:
                    solutions.add(info["solution"])

    grade = "상" if matched_1st else ("중" if matched_2nd else "하")
    return {
        "등급": grade,
        "카테고리": ", ".join(sorted(matched_cats)) or "-",
        "추천솔루션": ", ".join(sorted(solutions)) or "-",
        "매칭키워드": ", ".join(matched_1st + matched_2nd) or "-",
    }
