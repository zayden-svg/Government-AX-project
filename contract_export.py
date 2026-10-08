# contract_export.py
# 조달청 계약 전체 수집 → IT 관련만 저장 → 엑셀(진행중 / 곧 완료 / 업체 검색 / 자사 관련) 생성
#
#  · 수집: 조달청 계약정보 API(용역·물품)를 '하루 단위'로 최근 → 과거 순으로 끝까지 훑음.
#          하루 호출 한도가 있어 한 번에 다 못 끝내면 진행 위치를 저장하고 다음 실행 때 이어서 함.
#  · 저장: 정보화사업(Y) 또는 IT·자사 관련 단어가 사업명에 있는 계약만 g2b_contracts 테이블에 보관
#  · 엑셀: app_cache('contract_excel')에 저장 → 대시보드 낙찰결과 탭에서 내려받기
#
# 실행: python contract_export.py            (이어서 수집 + 엑셀 생성)
#       python contract_export.py --excel    (수집 없이 엑셀만 다시 생성)
#       python contract_export.py --out x.xlsx  (엑셀 파일도 로컬에 저장)
import argparse
import base64
import io
import json
import math
import os
import re
import sys
import time
from datetime import datetime, timedelta

import pandas as pd
from sqlalchemy import text

import common  # noqa: F401  (한국시간 고정)
from common import solution_hits, is_competitor_match
from db2 import get_engine
import store

TABLE = "g2b_contracts"
K_PROGRESS = "contract_crawl_progress"
K_EXCEL = "contract_excel"
K_EXCEL_META = "contract_excel_meta"   # 대시보드는 이것만 먼저 읽고, 파일은 버튼 누를 때 읽음
YEARS_BACK = 3                     # 과거 몇 년치 계약까지 훑을지 (진행 중인 계약은 대부분 3년 이내 체결)
CALL_BUDGET = int(os.getenv("CONTRACT_CALL_BUDGET", "900"))   # 한 번 실행에 쓸 최대 API 호출 수 (하루 한도 보호)
TIME_BUDGET_SEC = int(os.getenv("CONTRACT_TIME_BUDGET", str(80 * 60)))
ROWS_PER_PAGE = 999
SOON_DAYS = 90                     # '곧 완료' = 오늘부터 90일 안에 끝나는 계약
REORDER_LEAD_DAYS = 60
RECENT_REFRESH_DAYS = 7            # 매 실행마다 최근 7일은 다시 훑어 새 계약·변경 반영

OPS = {
    "용역": "https://apis.data.go.kr/1230000/ao/CntrctInfoService/getCntrctInfoListServcPPSSrch",
    "물품": "https://apis.data.go.kr/1230000/ao/CntrctInfoService/getCntrctInfoListThngPPSSrch",
}

IT_WORDS = [
    "정보시스템", "정보화", "시스템", "홈페이지", "누리집", "포털", "플랫폼", "소프트웨어", "SW", "S/W",
    "클라우드", "보안", "데이터", "AI", "인공지능", "전산", "서버", "네트워크", "웹", "앱", "모바일",
    "예약", "수강신청", "대기열", "트래픽", "부하테스트", "부하시험", "성능시험", "ISP", "ISMP", "차세대",
    "라이선스", "라이센스", "유지관리", "DB", "빅데이터", "디지털", "전자", "온라인", "통합관리",
]
_NON_IT = re.compile(r"(냉난방|공조|소방|승강기|조경|청소|경비|급식|방역|건축|토목|전기공사|도장|배관)")

# API 원본 필드 → 저장 칼럼 (가능한 한 많이 보관)
FIELD_MAP = [
    ("cntrct_no", "untyCntrctNo"), ("dcsn_no", "dcsnCntrctNo"), ("ref_no", "cntrctRefNo"),
    ("title", "cntrctNm"), ("cntrct_date", "cntrctCnclsDate"), ("start_date", "wbgnDate"),
    ("thtm_end", "thtmScmpltDate"), ("total_end", "ttalScmpltDate"), ("period", "cntrctPrd"),
    ("amount", "totCntrctAmt"), ("thtm_amount", "thtmCntrctAmt"),
    ("method", "cntrctCnclsMthdNm"), ("longterm", "lngtrmCtnuDivNm"), ("joint", "cmmnCntrctYn"),
    ("info_biz", "infoBizYn"), ("instt", "cntrctInsttNm"), ("instt_dept", "cntrctInsttChrgDeptNm"),
    ("instt_ofcl", "cntrctInsttOfclNm"), ("instt_tel", "cntrctInsttOfclTelNo"),
    ("dminstt_raw", "dminsttList"), ("corp_raw", "corpList"),
    ("ntce_no", "ntceNo"), ("req_no", "reqNo"), ("law", "baseLawNm"), ("pay", "payDivNm"),
    ("clsfc_l", "pubPrcrmntLrgclsfcNm"), ("clsfc_m", "pubPrcrmntMidclsfcNm"), ("clsfc", "pubPrcrmntClsfcNm"),
    ("rgst_dt", "rgstDt"), ("chg_dt", "chgDt"), ("url", "cntrctDtlInfoUrl"), ("info_url", "cntrctInfoUrl"),
]
COLS = ["uniq_key", "biz_type", "end_date", "end_est", "it_reason"] + [c for c, _ in FIELD_MAP] + ["updated_at"]


# ------------------------------------------------------------
# 저장소
# ------------------------------------------------------------
def ensure_table():
    with get_engine().begin() as conn:
        cols = ", ".join(f"{c} TEXT" for c in COLS if c != "uniq_key")
        conn.execute(text(f"CREATE TABLE IF NOT EXISTS {TABLE} (uniq_key TEXT PRIMARY KEY, {cols})"))


def save_rows(rows):
    if not rows:
        return 0
    ensure_table()
    dedup = {r["uniq_key"]: r for r in rows}
    cols = ", ".join(COLS)
    vals = ", ".join(f":{c}" for c in COLS)
    upd = ", ".join(f"{c}=excluded.{c}" for c in COLS if c != "uniq_key")
    stmt = text(f"INSERT INTO {TABLE} ({cols}) VALUES ({vals}) ON CONFLICT (uniq_key) DO UPDATE SET {upd}")
    data = [{c: str(r.get(c) or "") for c in COLS} for r in dedup.values()]
    with get_engine().begin() as conn:
        for i in range(0, len(data), 1000):
            conn.execute(stmt, data[i:i + 1000])
    return len(data)


def load_rows():
    ensure_table()
    with get_engine().begin() as conn:
        return pd.read_sql(text(f"SELECT * FROM {TABLE}"), conn)


# ------------------------------------------------------------
# 원본 → 행 변환
# ------------------------------------------------------------
def _date(v):
    s = re.sub(r"[^0-9]", "", str(v or ""))[:8]
    return f"{s[:4]}-{s[4:6]}-{s[6:8]}" if len(s) == 8 else ""


def corp_names(corp_raw):
    out = []
    for p in re.findall(r"\[([^\]]*)\]", str(corp_raw or "")):
        f = p.split("^")
        if len(f) >= 4 and f[3].strip():
            out.append(f[3].strip())
    return out


def dminstt_names(raw, fallback=""):
    out = []
    for p in re.findall(r"\[([^\]]*)\]", str(raw or "")):
        f = p.split("^")
        if len(f) >= 3 and f[2].strip():
            out.append(f[2].strip())
    return out or ([fallback] if fallback else [])


def _end_date(it):
    from collectors import _contract_end_date
    return _contract_end_date(it)


def it_reason(it):
    """IT 관련 판정 근거 (빈 문자열이면 IT 아님)"""
    title = str(it.get("cntrctNm") or "")
    if solution_hits(title):
        return "자사 관련"
    if str(it.get("infoBizYn") or "").upper() == "Y":
        return "정보화사업"
    if _NON_IT.search(title):
        return ""
    hit = next((w for w in IT_WORDS if w in title), "")
    return f"사업명({hit})" if hit else ""


def to_row(it, biz_type):
    reason = it_reason(it)
    if not reason:
        return None
    r = {c: str(it.get(k) or "").strip() for c, k in FIELD_MAP}
    for c in ("cntrct_date", "start_date", "thtm_end", "total_end"):
        r[c] = _date(r[c])
    end = _end_date(it)
    r["end_date"] = end or ""
    r["end_est"] = "" if end else "Y"
    r["biz_type"] = biz_type
    r["it_reason"] = reason
    r["uniq_key"] = f"{biz_type}|{r['cntrct_no'] or r['dcsn_no'] or r['title'] + r['cntrct_date']}"
    r["updated_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    return r


# ------------------------------------------------------------
# 수집 (하루 단위, 최근 → 과거, 이어하기)
# ------------------------------------------------------------
class QuotaExceeded(Exception):
    pass


class Crawler:
    def __init__(self, call_budget=CALL_BUDGET, time_budget=TIME_BUDGET_SEC):
        from collectors import _g2b_key
        self.key = _g2b_key()
        self.calls = 0
        self.call_budget = call_budget
        self.deadline = time.time() + time_budget

    def out_of_budget(self):
        return self.calls >= self.call_budget or time.time() > self.deadline

    def _page(self, url, day, page_no):
        from collectors import _g2b_get
        q = {"serviceKey": self.key, "pageNo": str(page_no), "numOfRows": str(ROWS_PER_PAGE), "type": "json",
             "inqryDiv": "1", "inqryBgnDate": day, "inqryEndDate": day}
        self.calls += 1
        resp = _g2b_get(url, q)
        try:
            data = resp.json()
        except ValueError:
            body = resp.text or ""
            if "LIMITED_NUMBER" in body or "22" in re.findall(r"<returnReasonCode>(\d+)<", body):
                raise QuotaExceeded("하루 호출 한도 초과")
            m = re.search(r"<(?:returnAuthMsg|errMsg|resultMsg)>([^<]+)<", body)
            raise RuntimeError(m.group(1) if m else body[:150])
        root = data.get("response") or {}
        header = root.get("header") or {}
        code = str(header.get("resultCode") or "00")
        if code == "22":
            raise QuotaExceeded(header.get("resultMsg") or "하루 호출 한도 초과")
        if code not in ("00", "0"):
            raise RuntimeError(f"{code} {header.get('resultMsg')}")
        body = root.get("body") or {}
        items = body.get("items") or []
        if isinstance(items, dict):
            items = items.get("item") or []
        if isinstance(items, dict):
            items = [items]
        total = int(body.get("totalCount") or len(items) or 0)
        return items, total

    def day(self, biz_type, day):
        """하루치 전체 페이지. 반환: (IT 행 목록, 원본 건수) — 예산이 모자라면 None"""
        url = OPS[biz_type]
        items, total = self._page(url, day, 1)
        per = max(len(items), 1)
        pages = math.ceil(total / per) if total else 1
        if self.calls + pages - 1 > self.call_budget:
            return None
        allit = list(items)
        for p in range(2, pages + 1):
            more, _ = self._page(url, day, p)
            if not more:
                break
            allit.extend(more)
        rows = [r for r in (to_row(it, biz_type) for it in allit) if r]
        return rows, len(allit)


def crawl():
    """진행 위치(유형별 '다음에 볼 날짜')에서 이어서 과거로 내려감. 최근 7일은 매번 다시 확인."""
    cr = Crawler()
    if not cr.key:
        print("[SKIP] G2B_SERVICE_KEY 없음 → 계약 수집 건너뜀")
        return {}
    prog, _ = store.load_cache(K_PROGRESS)
    prog = prog if isinstance(prog, dict) else {}
    today = datetime.now().date()
    floor = (today - timedelta(days=365 * YEARS_BACK)).strftime("%Y%m%d")
    stats = {"calls": 0, "saved": 0, "raw": 0, "stopped": ""}

    def _run_day(bt, d):
        res = cr.day(bt, d)
        if res is None:
            return False
        rows, raw = res
        stats["saved"] += save_rows(rows)
        stats["raw"] += raw
        return True

    try:
        # ① 최근 7일 새로고침
        for bt in OPS:
            for i in range(RECENT_REFRESH_DAYS):
                if cr.out_of_budget():
                    raise StopIteration
                _run_day(bt, (today - timedelta(days=i)).strftime("%Y%m%d"))
        # ② 과거로 이어서 (용역 먼저 끝까지 → 물품)
        for bt in OPS:
            st_ = prog.get(bt) or {}
            nxt = st_.get("next") or (today - timedelta(days=RECENT_REFRESH_DAYS)).strftime("%Y%m%d")
            while nxt >= floor:
                if cr.out_of_budget():
                    raise StopIteration
                if not _run_day(bt, nxt):
                    raise StopIteration
                nxt = (datetime.strptime(nxt, "%Y%m%d") - timedelta(days=1)).strftime("%Y%m%d")
                prog[bt] = {"next": nxt, "done": nxt < floor, "at": datetime.now().strftime("%Y-%m-%d %H:%M")}
                if cr.calls % 50 == 0:
                    store.save_cache(K_PROGRESS, prog)
            prog[bt] = {"next": nxt, "done": True, "at": datetime.now().strftime("%Y-%m-%d %H:%M")}
    except StopIteration:
        stats["stopped"] = "이번 실행 호출 한도·시간 도달 → 다음 실행에서 이어서"
    except QuotaExceeded as e:
        stats["stopped"] = f"조달청 하루 호출 한도 초과({e}) → 다음 실행에서 이어서"
    except Exception as e:
        stats["stopped"] = f"오류로 중단: {type(e).__name__}: {str(e)[:150]}"
    stats["calls"] = cr.calls
    prog["_last"] = {**stats, "at": datetime.now().strftime("%Y-%m-%d %H:%M")}
    store.save_cache(K_PROGRESS, prog)
    print(f"[OK] 계약 수집: API {cr.calls}회 · 원본 {stats['raw']}건 → IT 관련 {stats['saved']}건 저장 · "
          + json.dumps({k: v for k, v in prog.items() if k != "_last"}, ensure_ascii=False)
          + (f" · {stats['stopped']}" if stats["stopped"] else ""))
    return prog


# ------------------------------------------------------------
# 엑셀
# ------------------------------------------------------------
def _won(v):
    try:
        return int(float(str(v).replace(",", "")))
    except (TypeError, ValueError):
        return None


def prepare(df, competitors=None, today=None):
    today = pd.Timestamp(today or datetime.now().date())
    df = df.copy().fillna("")
    df["업체목록"] = df["corp_raw"].map(corp_names)
    df["업체명"] = df["업체목록"].map(lambda xs: " / ".join(xs))
    df["대표업체"] = df["업체목록"].map(lambda xs: xs[0] if xs else "")
    df["수요기관"] = [", ".join(dminstt_names(a, b)) for a, b in zip(df["dminstt_raw"], df["instt"])]
    end = pd.to_datetime(df["end_date"], errors="coerce")
    days = (end - today).dt.days
    df["남은일수"] = days
    df["상태"] = "종료일 미상"
    df.loc[days < 0, "상태"] = "완료"
    df.loc[(days >= 0) & (days <= SOON_DAYS), "상태"] = "곧 완료"
    df.loc[days > SOON_DAYS, "상태"] = "진행중"
    df["재발주예상"] = (end - pd.Timedelta(days=REORDER_LEAD_DAYS)).dt.strftime("%Y-%m-%d").fillna("")
    hits = df["title"].map(lambda t: solution_hits(t))
    df["자사관련단어"] = hits.map(lambda h: ", ".join(dict.fromkeys(h)) if h else "")
    df["자사관련"] = (df["자사관련단어"] != "").map({True: "Y", False: ""})
    comp = competitors or store.load_competitors()[0]
    df["경쟁사"] = df["업체명"].map(lambda c: "Y" if is_competitor_match(c, comp) else "")
    df["계약금액"] = df["amount"].map(_won)
    df["금차금액"] = df["thtm_amount"].map(_won)
    return df


OUT_COLS = [  # (엑셀 제목, 칼럼, 너비)
    ("상태", "상태", 9), ("남은일수", "남은일수", 8), ("구분", "biz_type", 6), ("계약명", "title", 46),
    ("업체명", "업체명", 26), ("수요기관", "수요기관", 24), ("계약금액(원)", "계약금액", 14),
    ("계약일", "cntrct_date", 11), ("착수일", "start_date", 11), ("종료일", "end_date", 11),
    ("종료일 추정", "end_est", 8), ("재발주 예상", "재발주예상", 11), ("자사관련", "자사관련", 7),
    ("자사관련 단어", "자사관련단어", 14), ("경쟁사", "경쟁사", 7), ("IT 판정", "it_reason", 12),
    ("계약방법", "method", 14), ("장기계속", "longterm", 9), ("공동계약", "joint", 7),
    ("금차금액(원)", "금차금액", 14), ("계약기간(원문)", "period", 20), ("계약기관", "instt", 20),
    ("담당부서", "instt_dept", 16), ("담당자", "instt_ofcl", 9), ("연락처", "instt_tel", 13),
    ("공고번호", "ntce_no", 14), ("요청번호", "req_no", 14), ("통합계약번호", "cntrct_no", 16),
    ("근거법령", "law", 14), ("분류", "clsfc", 14), ("등록일시", "rgst_dt", 16), ("원문", "url", 12),
]
STATUS_ORDER = {"곧 완료": 0, "진행중": 1, "종료일 미상": 2, "완료": 3}


def _sorted(df):
    d = df.assign(_o=df["상태"].map(STATUS_ORDER), _d=df["남은일수"].fillna(99999))
    return d.sort_values(["_o", "_d", "계약금액"], ascending=[True, True, False]).drop(columns=["_o", "_d"])


def build_excel(df, today=None):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    today_s = (pd.Timestamp(today) if today else pd.Timestamp(datetime.now())).strftime("%Y-%m-%d")
    wb = Workbook(write_only=True)
    navy = PatternFill("solid", fgColor="1B2A4A")
    hfont = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)
    big = Font(bold=True, size=16)
    soon_fill = PatternFill("solid", fgColor="FDECEA")
    mine_fill = PatternFill("solid", fgColor="E8F5E9")

    def cell(ws, v, font=None, fill=None, fmt=None, wrap=False):
        c = WriteOnlyCell(ws, value=v)
        if font:
            c.font = font
        if fill:
            c.fill = fill
        if fmt:
            c.number_format = fmt
        if wrap:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        return c

    def header(ws, titles):
        ws.append([cell(ws, t, hfont, navy) for t in titles])

    def table_sheet(name, sub, helper=False):
        ws = wb.create_sheet(name)
        for i, (_, _, w) in enumerate(OUT_COLS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "E2"
        titles = [t for t, _, _ in OUT_COLS] + (["검색키(수정 금지)"] if helper else [])
        header(ws, titles)
        name_col = get_column_letter([c for _, c, _ in OUT_COLS].index("업체명") + 1)
        for rn, r in enumerate(sub.to_dict("records"), start=2):
            row = []
            fill = soon_fill if r["상태"] == "곧 완료" else (mine_fill if r["자사관련"] == "Y" else None)
            for t, c, _ in OUT_COLS:
                v = r.get(c, "")
                if c == "남은일수":
                    v = None if pd.isna(v) else int(v)
                elif c in ("계약금액", "금차금액"):
                    v = None if v is None or (isinstance(v, float) and pd.isna(v)) else int(v)
                    row.append(cell(ws, v, fmt="#,##0"))
                    continue
                elif c == "url":
                    v = f'=HYPERLINK("{v}","열기")' if str(v).startswith("http") else ""
                row.append(cell(ws, v, fill=fill if c in ("상태", "title") else None))
            if helper:
                row.append(f'=IF(AND(업체검색!$B$3<>"",ISNUMBER(SEARCH(업체검색!$B$3,{name_col}{rn}))),ROW(),"")')
            ws.append(row)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(titles))}{max(len(sub) + 1, 2)}"
        return ws

    d = _sorted(df)
    n = d["상태"].value_counts()
    mine = d[(d["자사관련"] == "Y") | (d["경쟁사"] == "Y")]

    # 1) 안내·요약
    ws = wb.create_sheet("안내")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 70
    ws.append([cell(ws, "조달청 계약 현황 (IT 관련)", big)])
    ws.append([cell(ws, "기준일"), today_s])
    ws.append([cell(ws, "출처"), "조달청 나라장터 계약정보 OpenAPI (data.go.kr) — 용역·물품 계약, 최근 3년"])
    ws.append([cell(ws, "수록 범위"), "정보화사업(Y) 표시 계약 + 사업명에 IT·자사 관련 단어가 있는 계약 (건물·청소 등 비IT 제외)"])
    ws.append([])
    ws.append([cell(ws, "구분", hfont, navy), cell(ws, "건수 · 기준", hfont, navy)])
    for k, desc in [("곧 완료", f"오늘~{SOON_DAYS}일 안에 계약 종료 → 재발주 영업 대상"),
                    ("진행중", f"종료일이 {SOON_DAYS}일보다 더 남음"),
                    ("완료", "종료일이 지남"), ("종료일 미상", "계약기간 정보가 없음")]:
        ws.append([cell(ws, k, bold), f"{int(n.get(k, 0)):,}건 — {desc}"])
    ws.append([cell(ws, "자사 관련", bold), f"{int((d['자사관련'] == 'Y').sum()):,}건 — 대기열·접속·예약·매크로·부하테스트 등 자사 제품 관련 단어"])
    ws.append([cell(ws, "경쟁사 수주", bold), f"{int((d['경쟁사'] == 'Y').sum()):,}건 — 대시보드 공용 경쟁사 키워드 기준"])
    ws.append([cell(ws, "전체", bold), f"{len(d):,}건"])
    ws.append([])
    ws.append([cell(ws, "시트 안내", hfont, navy), cell(ws, "", hfont, navy)])
    for k, v in [("곧 완료", "90일 안에 끝나는 계약 (종료 임박 순)"), ("진행중", "현재 수행 중인 계약"),
                 ("자사 관련", "자사 제품 관련 단어 또는 경쟁사 수주 계약"),
                 ("업체별 요약", "업체마다 계약 수·진행중·곧 완료·금액 합계"),
                 ("업체검색", "B3 칸에 업체명 일부를 입력하면 전체 계약에서 찾아 보여줌 (최대 500건)"),
                 ("전체", "모든 계약. 머리글 ▼ 필터로 업체·기관·상태별 검색 가능")]:
        ws.append([cell(ws, k, bold), v])
    ws.append([])
    ws.append([cell(ws, "참고", bold), "종료일 = 총완수일 → 금차완수일 → 계약기간 문구 순으로 판단. 장기계속계약은 총완수일 기준."])
    ws.append([cell(ws, ""), f"재발주 예상 = 종료일 − {REORDER_LEAD_DAYS}일 (공공기관 통상 1~3개월 전 공고). 확정 정보 아님."])

    # 2) 업체 검색 (모든 엑셀 버전에서 동작하는 INDEX·SMALL 방식)
    ws = wb.create_sheet("업체검색")
    show = [("상태", "A"), ("남은일수", "B"), ("구분", "C"), ("계약명", "D"), ("업체명", "E"), ("수요기관", "F"),
            ("계약금액(원)", "G"), ("계약일", "H"), ("종료일", "J"), ("재발주 예상", "L"), ("자사관련", "M"), ("경쟁사", "O")]
    for i, w in enumerate([9, 8, 6, 46, 26, 24, 14, 11, 11, 11, 7, 7], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    helper_col = get_column_letter(len(OUT_COLS) + 1)
    ws.append([cell(ws, "업체 검색", big)])
    ws.append(["아래 노란 칸(B3)에 업체명 일부를 입력하세요 (예: 다이나패스). 전체 시트에서 찾아 최대 500건을 보여줍니다."])
    ws.append([cell(ws, "업체명 →", bold), cell(ws, "", fill=PatternFill("solid", fgColor="FFF59D")),
               cell(ws, "결과 건수", bold), f'=IF(B3="","",COUNT(전체!{helper_col}:{helper_col}))'])
    header(ws, [t for t, _ in show])
    for k in range(1, 501):
        rowf = []
        for t, col in show:
            tail = "" if t in ("남은일수", "계약금액(원)") else '&""'   # 빈 칸이 0으로 보이지 않게
            f = f'=IFERROR(INDEX(전체!{col}:{col},SMALL(전체!${helper_col}:${helper_col},{k})){tail},"")'
            rowf.append(cell(ws, f, fmt="#,##0") if t == "계약금액(원)" else f)
        ws.append(rowf)

    # 3) 시트별 표
    table_sheet("곧 완료", d[d["상태"] == "곧 완료"])
    table_sheet("진행중", d[d["상태"] == "진행중"])
    table_sheet("자사 관련", mine)

    # 4) 업체별 요약
    ws = wb.create_sheet("업체별 요약")
    ex = d.explode("업체목록")
    ex = ex[ex["업체목록"].fillna("") != ""]
    if not ex.empty:
        g = ex.groupby("업체목록")
        summ = pd.DataFrame({
            "계약 수": g.size(),
            "진행중": g["상태"].apply(lambda s: int((s == "진행중").sum())),
            "곧 완료": g["상태"].apply(lambda s: int((s == "곧 완료").sum())),
            "자사 관련": g["자사관련"].apply(lambda s: int((s == "Y").sum())),
            "금액 합계(원)": g["계약금액"].sum(min_count=1),
            "최근 계약일": g["cntrct_date"].max(),
            "주요 기관": g["수요기관"].apply(lambda s: ", ".join(s.value_counts().index[:3])),
            "경쟁사": g["경쟁사"].apply(lambda s: "Y" if (s == "Y").any() else ""),
        }).sort_values(["진행중", "계약 수"], ascending=False)
    else:
        summ = pd.DataFrame(columns=["계약 수", "진행중", "곧 완료", "자사 관련", "금액 합계(원)", "최근 계약일", "주요 기관", "경쟁사"])
    for i, w in enumerate([30, 8, 8, 8, 8, 16, 11, 50, 7], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    header(ws, ["업체명"] + list(summ.columns))
    for name, r in summ.iterrows():
        amt = r["금액 합계(원)"]
        ws.append([name, int(r["계약 수"]), r["진행중"], r["곧 완료"], r["자사 관련"],
                   cell(ws, None if pd.isna(amt) else int(amt), fmt="#,##0"), r["최근 계약일"], r["주요 기관"], r["경쟁사"]])
    ws.auto_filter.ref = f"A1:I{max(len(summ) + 1, 2)}"

    # 5) 전체 (검색키 포함)
    table_sheet("전체", d, helper=True)

    buf = io.BytesIO()
    wb.save(buf)
    meta = {"rows": len(d), "soon": int(n.get("곧 완료", 0)), "ongoing": int(n.get("진행중", 0)),
            "done": int(n.get("완료", 0)), "mine": len(mine), "companies": int(len(summ)), "date": today_s}
    return buf.getvalue(), meta


def export_excel(out_path=None):
    df = load_rows()
    if df.empty:
        print("[SKIP] 저장된 계약이 없어 엑셀을 만들지 않음")
        return None
    data, meta = build_excel(prepare(df))
    fname = f"조달청_계약현황_{meta['date'].replace('-', '')}.xlsx"
    store.save_cache(K_EXCEL, {"filename": fname, "b64": base64.b64encode(data).decode(), **meta,
                               "size": len(data)})
    store.save_cache(K_EXCEL_META, {"filename": fname, **meta, "size": len(data)})
    if out_path:
        with open(out_path, "wb") as f:
            f.write(data)
    print(f"[OK] 계약 엑셀: {fname} · {len(data) / 1e6:.1f}MB · {json.dumps(meta, ensure_ascii=False)}")
    return meta


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", action="store_true", help="수집 없이 엑셀만 생성")
    ap.add_argument("--out", default="", help="엑셀 파일 로컬 저장 경로")
    a = ap.parse_args()
    if not a.excel:
        crawl()
    export_excel(a.out or None)


if __name__ == "__main__":
    sys.exit(main())
