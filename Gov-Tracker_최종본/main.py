# main.py — 매일 아침 자동수집: 수집 → AI 분석 → 저장 → 낙찰·계약 결과
import os
import sys
import traceback

import common  # noqa: F401  (한국시간 고정 — 다른 모듈보다 먼저)
import asyncio
import hashlib
import re
from datetime import datetime, date

import pandas as pd
from sqlalchemy import text, bindparam
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from collectors import run_all_collectors, fetch_g2b_results
from common import is_mois_noise
from store import load_cache, save_cache
from procurement_store import save_results, count_results
from biz_classifier import classify_and_score as biz_classify_and_score
from ai_utils import is_ai_ready, analyze_postings_batch, AI_MAX_CONCURRENCY, AI_RPM_LIMIT, MODEL_NAME
from db2 import get_engine, is_postgres

EXCEL_OUTPUT = "Gov-Tracker_결과.xlsx"
COLLECT_LIMIT = int(os.getenv("COLLECT_LIMIT", "20"))  # 수집처당 최대 건수

# ------------------------------------------------------------
# 공고 유형(입찰/RND/기타) - 이건 "AI 구분"과 다른 개념.
# 게시글 형식이 입찰공고문인지 R&D공모문인지만 가볍게 나누는 용도.
# ------------------------------------------------------------
RND_TYPE_KEYWORDS = ["과제공고", "지원사업", "R&D", "연구개발", "공모전", "지원과제", "사업 공모", "공모", "수요조사"]
BID_TYPE_KEYWORDS = ["입찰", "전자입찰", "구매", "용역", "발주", "제안서", "제안", "적격심사", "사전규격"]

TRACK_LABELS = {"RND": "R&D", "BIZ": "사업부", "": "미분류"}

# AI 분석이 실패했을 때만 쓰는 최후의 안전장치 (평소엔 사용 안 함)
AGENCY_TRACK_FALLBACK = {
    "IRIS": "RND", "KERIS": "RND", "NTIS": "RND", "TIPA": "RND",
    "KIAT": "RND", "INNOPOLIS": "RND", "국가AI전략위원회": "RND", "AIHub": "RND",
    "NIPA": "BIZ", "조달청": "BIZ", "조달청(사전규격)": "BIZ", "행정안전부": "BIZ",
}


def classify_post_type(title: str) -> str:
    if not title:
        return "ETC"
    for kw in BID_TYPE_KEYWORDS:
        if kw in title:
            return "BID"
    for kw in RND_TYPE_KEYWORDS:
        if kw in title:
            return "RND"
    return "ETC"


# ------------------------------------------------------------
# 담당자명 정제
# ------------------------------------------------------------
AGENCY_SUFFIXES = [
    "센터", "재단", "진흥원", "위원회", "연구원", "정보원",
    "공사", "청", "부", "원", "팀", "실", "국", "협회", "본부",
]


def clean_manager_name(raw) -> str:
    if raw is None:
        return ""
    name = str(raw).strip()
    if not name:
        return ""
    if re.search(r"\d", name):
        return ""
    for suffix in AGENCY_SUFFIXES:
        if name.endswith(suffix):
            return ""
    if re.fullmatch(r"[가-힣]{2,4}", name):
        return name
    if re.fullmatch(r"[A-Za-z]{2,20}(\s[A-Za-z]{2,20})?", name):
        return name
    return ""


# ------------------------------------------------------------
# 중복 판별용 해시 (기관 + 제목 + 등록일 + 마감일 기준)
# ------------------------------------------------------------
def normalize_for_dedup(text_val) -> str:
    if not text_val:
        return ""
    t = str(text_val)
    t = re.sub(r"\s+", "", t)
    t = re.sub(r"[^\w가-힣]", "", t)
    return t.lower()


def build_dedup_hash(title: str, agency: str, reg_date: str, due_date: str) -> str:
    norm_title = normalize_for_dedup(title)
    norm_agency = normalize_for_dedup(agency)
    base = f"{norm_agency}|{norm_title}|{reg_date or ''}|{due_date or ''}"
    return hashlib.md5(base.encode("utf-8")).hexdigest()


def make_uniq_key(rec: dict) -> str:
    base = f"{rec.get('agency','')}|{rec.get('title','')}|{rec.get('reg_date','')}|{rec.get('url','')}"
    return hashlib.md5(base.encode("utf-8")).hexdigest()


# ------------------------------------------------------------
# DB 초기화 / 마이그레이션 (Supabase-Postgres, 로컬-SQLite 겸용)
# ------------------------------------------------------------
def init_db():
    engine = get_engine()
    with engine.begin() as conn:
        conn.execute(text("""
            CREATE TABLE IF NOT EXISTS postings (
                uniq_key TEXT PRIMARY KEY,
                source TEXT, agency TEXT, gubun TEXT, post_type TEXT, title TEXT,
                dept TEXT, manager TEXT, reg_date TEXT, due_date TEXT, budget TEXT,
                attach TEXT, views TEXT, url TEXT, grade TEXT, category TEXT,
                matched_keywords TEXT, recommended_solution TEXT, status TEXT,
                track TEXT, track_reason TEXT,
                ai_priority_score INTEGER, ai_priority_reason TEXT,
                dedup_hash TEXT, ai_oneline TEXT, ai_summary TEXT, content TEXT,
                created_at TEXT, updated_at TEXT
            )
        """))
    ensure_columns(engine)
    with engine.begin() as conn:
        conn.execute(text("CREATE INDEX IF NOT EXISTS idx_dedup_hash ON postings(dedup_hash)"))
    return engine


def ensure_columns(engine):
    """예전 버전 DB에 새 컬럼(track_reason, ai_priority_score 등)이 없으면 추가"""
    needed = {c: "TEXT" for c in [
        "source", "agency", "gubun", "post_type", "title", "dept", "manager", "reg_date", "due_date",
        "budget", "attach", "views", "url", "grade", "category", "matched_keywords", "recommended_solution",
        "status", "track", "track_reason", "ai_priority_reason", "dedup_hash", "ai_oneline", "ai_summary",
        "content", "created_at", "updated_at",
    ]}
    needed["ai_priority_score"] = "INTEGER"
    with engine.begin() as conn:
        if is_postgres():
            existing = {row[0] for row in conn.execute(text(
                "SELECT column_name FROM information_schema.columns WHERE table_name='postings'"
            ))}
        else:
            existing = {row[1] for row in conn.execute(text("PRAGMA table_info(postings)"))}
        for col, coltype in needed.items():
            if col not in existing:
                conn.execute(text(f"ALTER TABLE postings ADD COLUMN {col} {coltype}"))


def load_existing_index(engine):
    """기존 레코드를 DB에서 한 번에 읽어 메모리 사전으로 만든다 (공고 1건마다 DB를 왕복하지 않음)."""
    by_hash, by_key = {}, {}
    with engine.begin() as conn:
        rows = conn.execute(text("""
            SELECT uniq_key, dedup_hash, track, track_reason, ai_priority_score,
                   ai_priority_reason, ai_oneline, ai_summary, created_at, budget
            FROM postings
        """)).fetchall()
    for r in rows:
        by_key[r[0]] = r
        if r[1] and r[1] not in by_hash:
            by_hash[r[1]] = r
    return by_hash, by_key


_REC_FIELDS = [
    "source", "agency", "gubun", "post_type", "title", "dept", "manager", "reg_date",
    "due_date", "budget", "attach", "views", "url", "grade", "category", "matched_keywords",
    "recommended_solution", "track", "track_reason", "ai_priority_reason", "dedup_hash",
    "ai_oneline", "ai_summary", "content",
]

INSERT_SQL = text("""
    INSERT INTO postings (
        uniq_key, source, agency, gubun, post_type, title, dept, manager,
        reg_date, due_date, budget, attach, views, url,
        grade, category, matched_keywords, recommended_solution,
        track, track_reason, ai_priority_score, ai_priority_reason, dedup_hash,
        ai_oneline, ai_summary, content, status, created_at, updated_at
    ) VALUES (
        :uniq_key, :source, :agency, :gubun, :post_type, :title, :dept, :manager,
        :reg_date, :due_date, :budget, :attach, :views, :url,
        :grade, :category, :matched_keywords, :recommended_solution,
        :track, :track_reason, :ai_priority_score, :ai_priority_reason, :dedup_hash,
        :ai_oneline, :ai_summary, :content, '진행중', :created_at, :now
    )
""")


def upsert_postings(engine, records, existing_by_key):
    """전체 레코드를 트랜잭션 1번으로 저장. 같은 키는 지우고 다시 넣되 최초 수집시각·기존 예산은 보존."""
    if not records:
        return
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    params = {}
    for rec in records:
        key = rec.get("_uniq_key") or make_uniq_key(rec)
        p = {f: str(rec.get(f) or "") for f in _REC_FIELDS}
        p["content"] = p["content"][:1000]
        p["uniq_key"] = key
        p["ai_priority_score"] = rec.get("ai_priority_score")
        old = existing_by_key.get(key)
        p["created_at"] = (old[8] if old and old[8] else now)
        if not p["budget"] and old and old[9]:
            p["budget"] = str(old[9])          # 이번에 예산을 못 찾았으면 예전에 찾은 값 유지
        p["now"] = now
        params[key] = p
    keys = list(params)
    with engine.begin() as conn:
        for i in range(0, len(keys), 500):
            stmt = text("DELETE FROM postings WHERE uniq_key IN :ks").bindparams(bindparam("ks", expanding=True))
            conn.execute(stmt, {"ks": keys[i:i + 500]})
        conn.execute(INSERT_SQL, list(params.values()))


def mark_expired(engine):
    today = date.today().strftime("%Y-%m-%d")
    with engine.begin() as conn:
        conn.execute(text("""
            UPDATE postings SET status='마감'
            WHERE due_date != '' AND due_date < :today AND status != '마감'
        """), {"today": today})


def _apply_grade_style(path):
    wb = load_workbook(path)
    ws = wb["전체"]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    orange_fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")

    headers = [cell.value for cell in ws[1]]
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    if "등급" in headers:
        grade_col = headers.index("등급") + 1
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            grade_cell = row[grade_col - 1]
            if grade_cell.value == "상":
                for c in row:
                    c.fill = red_fill
            elif grade_cell.value == "중":
                for c in row:
                    c.fill = orange_fill

    if "제목" in headers and "원문링크" in headers:
        title_col = headers.index("제목") + 1
        url_col = headers.index("원문링크") + 1
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            title_cell = row[title_col - 1]
            url_cell = row[url_col - 1]
            if url_cell.value:
                title_cell.hyperlink = url_cell.value
                title_cell.font = Font(color="0563C1", underline="single")
        ws.delete_cols(url_col)

    for col_cells in ws.columns:
        max_len = max((len(str(c.value)) if c.value else 0) for c in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)

    ws.freeze_panes = "A2"
    wb.save(path)


def export_to_excel(engine):
    df = pd.read_sql_query("""
        SELECT
            agency AS 기관, track AS 분류, track_reason AS AI구분근거,
            post_type AS 공고유형, gubun AS 세부구분, title AS 제목,
            dept AS 담당부서, manager AS 담당자, reg_date AS 등록일, due_date AS 마감일,
            budget AS 예산, ai_oneline AS 한줄요약,
            grade AS 등급, category AS 카테고리, recommended_solution AS 추천솔루션,
            matched_keywords AS 매칭키워드, ai_priority_score AS AI연관도점수,
            ai_priority_reason AS AI연관도근거, status AS 상태, url AS 원문링크
        FROM postings
        ORDER BY reg_date DESC
    """, engine)

    df = df.fillna("")
    df["분류"] = df["분류"].map(lambda v: TRACK_LABELS.get(v, v) if v else "미분류")

    with pd.ExcelWriter(EXCEL_OUTPUT, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="전체", index=False)

    _apply_grade_style(EXCEL_OUTPUT)
    print(f"[완료] 결과 파일 저장: {EXCEL_OUTPUT} ({len(df)}건)")


def _build_ai_info_block(rec: dict, agency: str) -> str:
    return (
        f"제목: {rec.get('title', '')}\n"
        f"기관: {rec.get('agency', agency)}\n"
        f"부서: {rec.get('dept', '')}\n"
        f"공고유형(원본 구분값): {rec.get('gubun', '')}\n"
        f"접수 마감일: {rec.get('due_date', '') or '미표기'}\n"
        f"예산(원): {rec.get('budget', '') or '미표기'}\n"
        f"본문 일부: {str(rec.get('content', ''))[:1000]}"
    )


def collect_and_process():
    engine = init_db()
    new_count = updated_count = skipped_count = duplicate_count = ai_called_count = 0

    results_by_agency = run_all_collectors(limit=COLLECT_LIMIT)
    total_records = sum(len(v) for v in results_by_agency.values())
    print(f"\n[1단계] 총 {total_records}건 수집 완료. 등급판정/중복확인 진행 중...\n")

    existing_by_hash, existing_by_key = load_existing_index(engine)
    seen_hashes = set()          # 이번 수집분 안에서의 중복 방지
    ready_records, pending_ai, rec_map = [], [], {}

    today_str = datetime.now().strftime("%Y-%m-%d")
    for agency, records in results_by_agency.items():
        for rec in records:
            if not rec.get("title"):
                skipped_count += 1
                continue
            # 이미 마감된 공고·일반 보도자료는 저장·AI 분석하지 않음 (비용 절감)
            due = str(rec.get("due_date") or "")
            if (len(due) == 10 and due < today_str) or is_mois_noise(rec.get("agency", agency), rec["title"]):
                skipped_count += 1
                continue

            rec["manager"] = clean_manager_name(rec.get("manager", ""))
            rec["post_type"] = classify_post_type(rec["title"])

            score_info = biz_classify_and_score(rec["title"], rec.get("content", ""))
            rec["grade"] = score_info["등급"]
            rec["category"] = score_info["카테고리"]
            rec["recommended_solution"] = score_info["추천솔루션"]
            rec["matched_keywords"] = score_info["매칭키워드"]

            dedup_hash = build_dedup_hash(
                rec.get("title", ""), rec.get("agency", agency),
                rec.get("reg_date", ""), rec.get("due_date", "")
            )
            rec["dedup_hash"] = dedup_hash
            key = make_uniq_key(rec)
            rec["_uniq_key"] = key

            if dedup_hash in seen_hashes:
                duplicate_count += 1
                continue
            seen_hashes.add(dedup_hash)

            existing = existing_by_hash.get(dedup_hash)
            if existing and existing[0] != key:
                duplicate_count += 1
                continue

            # 기존 AI 결과가 있고 요약까지 채워져 있으면 재사용 (AI 비용 0)
            if existing and existing[4] is not None and existing[7]:
                rec["track"], rec["track_reason"] = existing[2], existing[3]
                rec["ai_priority_score"], rec["ai_priority_reason"] = existing[4], existing[5]
                rec["ai_oneline"], rec["ai_summary"] = existing[6], existing[7]
                ready_records.append(rec)
            else:
                pending_ai.append((key, _build_ai_info_block(rec, agency)))
                rec_map[key] = rec

    ai_total = len(pending_ai)
    print(f"\n[2단계] 신규 AI 분석 {ai_total}건 (모델 {MODEL_NAME}, 동시 {AI_MAX_CONCURRENCY}건, 분당 {AI_RPM_LIMIT}건)\n")

    ai_results = {}
    if ai_total and is_ai_ready():
        done = {"n": 0}

        def _progress(k, result):
            done["n"] += 1
            title = rec_map[k].get("title", "")[:40]
            if result and not result.get("error"):
                print(f"  [{done['n']}/{ai_total}] {result['track']} / {result['score']}점 - {title}")
            else:
                print(f"  [{done['n']}/{ai_total}] 실패: {(result or {}).get('error', '알수없음')} - {title}")

        ai_results = asyncio.run(analyze_postings_batch(pending_ai, progress_cb=_progress))
        ai_called_count = ai_total
    elif ai_total:
        print("  [경고] ANTHROPIC_API_KEY 미설정 → AI 분석 생략, 기관 기본값으로 임시 분류")

    for k, rec in rec_map.items():
        r = ai_results.get(k)
        if r and not r.get("error"):
            rec["track"], rec["track_reason"] = r["track"], r["track_reason"]
            rec["ai_priority_score"], rec["ai_priority_reason"] = r["score"], r["score_reason"]
            rec["ai_oneline"], rec["ai_summary"] = r["oneline"], r["summary"]
        else:
            err_msg = (r or {}).get("error", "API 키 미설정")
            old = existing_by_key.get(k)
            if old and old[4] is not None:      # 예전 AI 결과가 있으면 지우지 않고 유지
                rec["track"], rec["track_reason"] = old[2], old[3]
                rec["ai_priority_score"], rec["ai_priority_reason"] = old[4], old[5]
                rec["ai_oneline"], rec["ai_summary"] = old[6] or "", old[7] or ""
            else:
                rec["track"] = AGENCY_TRACK_FALLBACK.get(rec.get("agency", ""), "BIZ")
                rec["track_reason"] = f"AI 분석 실패({err_msg})로 기관 기본값으로 임시 분류됨"
                rec["ai_priority_score"] = None
                rec["ai_priority_reason"] = "AI 분석 대기 중"
        ready_records.append(rec)

    print(f"\n[3단계] DB 저장 중... (총 {len(ready_records)}건)\n")
    for rec in ready_records:
        if rec["_uniq_key"] in existing_by_key:
            updated_count += 1
        else:
            new_count += 1
    upsert_postings(engine, ready_records, existing_by_key)
    mark_expired(engine)

    try:
        export_to_excel(engine)
    except Exception as e:                     # 엑셀은 부가 산출물 — 실패해도 수집 결과에는 영향 없음
        print(f"[경고] 엑셀 저장 실패(건너뜀): {e}")

    # [4단계] 조달청 낙찰·계약 결과 (시장 정보) — 실패해도 위 작업에는 영향 없음
    #   처음 실행(저장된 결과 50건 미만)이면 최근 1년치를 한 번 채워 재발주 알림에 활용
    try:
        done_flag, _ = load_cache("procurement_backfill_done")
        backfill = 365 if (not done_flag and count_results() < 50) else 0
        if backfill:
            print("[4단계] 낙찰·계약 최초 실행 — 최근 1년치 자사 관련 사업을 함께 조회합니다.")
        saved = save_results(fetch_g2b_results(backfill_days=backfill))
        print(f"[4단계] 낙찰·계약 결과 {saved}건 저장")
        if backfill and saved:
            save_cache("procurement_backfill_done", {"date": today_str, "saved": saved})
    except Exception as e:
        traceback.print_exc()
        print(f"[4단계] 낙찰·계약 수집 실패(건너뜀): {e}")

    print(
        f"\n신규 {new_count}건 / 갱신 {updated_count}건 / 제외(제목없음·마감·보도자료) {skipped_count}건 / "
        f"중복 제외 {duplicate_count}건 / AI 신규 분석 {ai_called_count}건 처리 완료"
    )


if __name__ == "__main__":
    collect_and_process()
    sys.exit(0)
