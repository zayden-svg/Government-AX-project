"""
Gov-Tracker 통합 수집 스크립트 (biz_classifier 연동 / 클릭캡처 링크수집판)
실행: python gov_tracker.py
대상 파일: Gov-Tracker_누적데이터.xlsx (같은 폴더에 있어야 함, 실행 전 Excel에서 닫아둘 것)
같은 폴더 필요 파일: biz_classifier.py
"""

import re
import time
from datetime import datetime, timedelta, date
from pathlib import Path

import requests
from playwright.sync_api import sync_playwright
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── biz_classifier 연결 ──────────────────────────────────────
from biz_classifier import classify_and_score

EXCEL_FILE = "Gov-Tracker_누적데이터.xlsx"
DASHBOARD_SHEET = "요약"
ARCHIVE_SHEET = "보관"

# ── 기관별 컬럼 구성 (등급/카테고리/추천솔루션/매칭키워드 4개 공통 추가) ──
CLASSIFY_COLUMNS = ["등급", "카테고리", "추천솔루션", "매칭키워드"]

SHEET_COLUMNS = {
    "조달청":       ["기관", "공고구분", "제목", "담당자", "등록일", "마감일", "첨부파일", "원문링크"] + CLASSIFY_COLUMNS,
    "알리오":       ["기관", "공고구분", "제목", "등록일", "마감일", "원문링크"] + CLASSIFY_COLUMNS,
    "행정안전부":   ["기관", "공고구분", "제목", "담당자", "등록일", "조회수", "첨부파일", "원문링크"] + CLASSIFY_COLUMNS,
    "NIPA":         ["기관", "공고구분", "제목", "담당자", "등록일", "조회수", "첨부파일", "원문링크"] + CLASSIFY_COLUMNS,
    "AIHub":        ["기관", "공고구분", "제목", "등록일", "조회수", "원문링크"] + CLASSIFY_COLUMNS,
    "국가AI전략위원회": ["기관", "공고구분", "제목", "등록일", "조회수", "원문링크"] + CLASSIFY_COLUMNS,
    "IITP":         ["기관", "공고구분", "제목", "담당자", "등록일", "마감일", "원문링크"] + CLASSIFY_COLUMNS,
    "NIA":          ["기관", "공고구분", "제목", "등록일", "조회수", "첨부파일", "원문링크"] + CLASSIFY_COLUMNS,
    "KERIS":        ["기관", "공고구분", "제목", "담당부서", "담당자", "등록일", "마감일", "원문링크"] + CLASSIFY_COLUMNS,
}
ARCHIVE_COLUMNS_BASE = ["기관", "공고구분", "제목", "등록일", "마감일", "원문링크", "원본시트", "보관일"]

COL_WIDTH_HINTS = {
    "기관": 24, "공고구분": 14, "제목": 55, "담당부서": 16, "담당자": 12,
    "등록일": 16, "마감일": 16, "조회수": 10, "첨부파일": 10, "원문링크": 45,
    "원본시트": 14, "보관일": 14,
    "등급": 6, "카테고리": 28, "추천솔루션": 18, "매칭키워드": 30,
}

# ── 서비스키 (조달청 공공데이터포털) ──────────────────────────
G2B_SERVICE_KEY = "719c800f7f4f184abf11b2c05252d7bd703176c3c040d316c63e5bd50918f848"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
}

ALIO_AGENCIES = {
    "한국인터넷진흥원": "한국인터넷진흥원(KISA)",
    "한국콘텐츠진흥원": "한국콘텐츠진흥원(KOCCA)",
    "정보통신정책연구원": "정보통신정책연구원(KISDI)",
}

# ── 등급별 색상 (biz_classifier 결과 기반) ───────────────────
GRADE_FILL = {
    "상": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),  # 붉은 계열
    "중": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),  # 노란 계열
    "하": PatternFill(fill_type=None),
}
GRADE_FONT = {
    "상": Font(bold=True, color="9C0006"),
    "중": Font(bold=False, color="9C6500"),
    "하": Font(),
}
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
URGENT_BORDER = Border(
    left=Side(style="medium", color="FF0000"), right=Side(style="medium", color="FF0000"),
    top=Side(style="medium", color="FF0000"), bottom=Side(style="medium", color="FF0000"),
)
NO_FILL = PatternFill(fill_type=None)
NO_BORDER = Border()

RENOTICE_PATTERNS = [
    r"\(재공고\)", r"\[재공고\]", r"재공고",
    r"\(변경공고\)", r"\[변경공고\]", r"변경공고",
    r"\(수정공고\)", r"\[수정공고\]", r"수정공고", r"\(수정\)", r"\[수정\]",
    r"\(연장공고\)", r"\[연장공고\]", r"연장공고",
    r"\(취소공고\)", r"\[취소공고\]", r"취소공고",
    r"\(긴급\)", r"\[긴급\]", r"긴급",
    r"\(재입찰공고\)", r"\[재입찰공고\]", r"재입찰공고",
    r"\(재입찰\)", r"\[재입찰\]",
]


def normalize_title(title: str) -> str:
    if not title:
        return ""
    t = str(title)
    for pat in RENOTICE_PATTERNS:
        t = re.sub(pat, "", t)
    return re.sub(r"\s+", " ", t).strip()


def parse_date(value):
    if value is None:
        return None
    s = str(value).strip()
    if not s or s.lower() == "nan":
        return None
    s = s.rstrip(".")
    fmts = (
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
        "%Y.%m.%d %H:%M:%S", "%Y.%m.%d %H:%M", "%Y.%m.%d",
        "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y/%m/%d",
    )
    for fmt in fmts:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    m = re.search(r"(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


def is_still_open(due_date_str) -> bool:
    d = parse_date(due_date_str)
    if d is None:
        return True
    return d >= date.today()


def filter_active(items):
    return [it for it in items if is_still_open(it.get("마감일", ""))]


def fix_link(link: str, base: str) -> str:
    """상대경로 링크를 절대경로로 보정. 이미 http(s)면 그대로 반환."""
    if not link:
        return ""
    link = link.strip()
    if link.lower().startswith("javascript") or link == "#":
        return ""
    if link.startswith("http"):
        return link
    if link.startswith("/"):
        return base + link
    return base + "/" + link


def click_and_capture_url(page, clickable_el, list_url, timeout=8000):
    """제목 요소를 클릭해서 실제로 이동하는 URL을 캡처한 뒤 목록으로 복귀.
    href/onclick을 추측하는 대신 브라우저가 실제로 가는 곳을 그대로 확인하는 방식."""
    try:
        with page.expect_navigation(timeout=timeout):
            clickable_el.click()
        real_url = page.url
        page.go_back(wait_until="networkidle", timeout=timeout)
        return real_url
    except Exception:
        try:
            page.goto(list_url, wait_until="networkidle", timeout=timeout)
        except Exception:
            pass
        return ""


# ============================================================
# 1. 조달청
# ============================================================
def fetch_g2b():
    results = []
    end = datetime.now()
    start = end - timedelta(days=14)
    bgn_str = start.strftime("%Y%m%d0000")
    end_str = end.strftime("%Y%m%d2359")

    bid_url = ("https://apis.data.go.kr/1230000/ad/BidPublicInfoService/"
               "getBidPblancListInfoServcPPSSrch")
    params = {
        "ServiceKey": G2B_SERVICE_KEY, "type": "json", "inqryDiv": "1",
        "inqryBgnDt": bgn_str, "inqryEndDt": end_str,
        "numOfRows": "100", "pageNo": "1",
    }
    try:
        r = requests.get(bid_url, params=params, headers=HEADERS, timeout=20)
        data = r.json()
        header = data.get("response", {}).get("header", {})
        if header.get("resultCode") == "00":
            items = data.get("response", {}).get("body", {}).get("items")
            if isinstance(items, dict):
                items = items.get("item", [])
            if isinstance(items, dict):
                items = [items]
            for it in items or []:
                results.append({
                    "기관": it.get("dminsttNm", "조달청"),
                    "공고구분": "입찰공고",
                    "제목": it.get("bidNtceNm", ""),
                    "담당자": it.get("ntceInsttOfclNm", ""),
                    "등록일": it.get("bidNtceDt", "") or it.get("rgstDt", ""),
                    "마감일": it.get("bidClseDt", ""),
                    "첨부파일": "Y" if it.get("ntceSpecDocUrl1") else "",
                    "원문링크": it.get("bidNtceDtlUrl", ""),
                })
        else:
            print(f"  [조달청-입찰] API 오류: {header.get('resultCode')} {header.get('resultMsg')}")
    except Exception as e:
        print(f"  [조달청-입찰] 수집 실패: {e}")

    prespec_url = ("https://apis.data.go.kr/1230000/ao/HrcspSsstndrdInfoService/"
                    "getPublicPrcureThngInfoServcPPSSrch")
    try:
        r = requests.get(prespec_url, params=params, headers=HEADERS, timeout=20)
        data = r.json()
        header = data.get("response", {}).get("header", {})
        if header.get("resultCode") == "00":
            items = data.get("response", {}).get("body", {}).get("items")
            if isinstance(items, dict):
                items = items.get("item", [])
            if isinstance(items, dict):
                items = [items]
            for it in items or []:
                bf_spec_no = it.get("bfSpecRgstNo", "")
                rcpt = it.get("rcptDt", "")
                reg_date = rcpt.split(" ")[0] if rcpt else ""
                results.append({
                    "기관": it.get("orderInsttNm", "조달청"),
                    "공고구분": "사전규격공고",
                    "제목": it.get("prdctClsfcNoNm", ""),
                    "담당자": it.get("ofclNm", ""),
                    "등록일": reg_date,
                    "마감일": it.get("opninRgstClseDt", ""),
                    "첨부파일": "Y" if it.get("specDocFileUrl1") else "",
                    "원문링크": f"https://www.g2b.go.kr/pn/pnz/pnza/bidPrgsInfo/openSpecView.do?bfSpecRegNo={bf_spec_no}" if bf_spec_no else "",
                })
        else:
            print(f"  [조달청-사전규격] API 오류: {header.get('resultCode')} {header.get('resultMsg')}")
    except Exception as e:
        print(f"  [조달청-사전규격] 수집 실패: {e}")

    return filter_active(results)


# ============================================================
# 2. 알리오
# ============================================================
def fetch_alio():
    results = []
    for keyword, display_name in ALIO_AGENCIES.items():
        url = "https://alio.go.kr/occasional/findBidList.json"
        params = {"type": "apbaNa", "word": keyword, "pageNo": "1"}
        for attempt in range(2):
            try:
                r = requests.get(url, params=params, headers=HEADERS, timeout=15)
                data = r.json()
                total_cnt = (data.get("data") or {}).get("totalCnt", 0)
                items = (data.get("data") or {}).get("result") or []
                if total_cnt == 0 or not items:
                    if attempt == 0:
                        time.sleep(2)
                        continue
                    print(f"  [알리오-{display_name}] 0건 (API 응답 totalCnt=0)")
                    break
                count = 0
                for it in items[:10] if isinstance(items, list) else []:
                    title = it.get("bidNm") or it.get("title") or ""
                    seq = it.get("seq") or it.get("bidSeq") or ""
                    reg_date = it.get("ntceDt") or it.get("rgstDt") or ""
                    due_date = it.get("bidClseDt") or it.get("clseDt") or ""
                    results.append({
                        "기관": display_name, "공고구분": "입찰공고", "제목": title,
                        "등록일": reg_date, "마감일": due_date,
                        "원문링크": f"https://alio.go.kr/occasional/bidDtl.do?seq={seq}&type=apbaNa&pageNo=1" if seq else "",
                    })
                    count += 1
                print(f"  -> [{display_name}] {count}건 수집")
                break
            except Exception as e:
                print(f"  [알리오-{display_name}] 수집 실패: {e}")
                break

    if not results:
        print("  [알리오] 전체 0건 — API 응답 구조가 바뀌었을 가능성. "
              "alio_debug.json 저장을 원하면 fetch_alio()에 r.text 저장 코드를 추가하세요.")
    return filter_active(results)


# ============================================================
# 3. NIPA (Playwright) — 열 순서 버그 수정: 번호|제목|작성자|파일|조회수|작성일
# ============================================================
def fetch_nipa(page):
    results = []
    url = "https://www.nipa.kr/home/2-3"
    try:
        page.goto(url, wait_until="networkidle", timeout=30000)
        page.wait_for_timeout(2000)
        try:
            page.wait_for_selector("table tbody tr", timeout=8000)
        except Exception:
            pass
        rows = page.query_selector_all("table tbody tr")
        for row in rows[:15]:
            cells = row.query_selector_all("td")
            if len(cells) < 5:
                continue
            title_el = cells[1].query_selector("a") or row.query_selector("a")
            title = title_el.inner_text().strip() if title_el else cells[1].inner_text().strip()
            link = title_el.get_attribute("href") if title_el else ""
            link = fix_link(link, "https://www.nipa.kr")

            # 실제 열 순서: [0]번호 [1]제목 [2]작성자 [3]파일 [4]조회수 [5]작성일
            author = cells[2].inner_text().strip() if len(cells) > 2 else ""
            views = cells[4].inner_text().strip() if len(cells) > 4 else ""
            reg_date = cells[5].inner_text().strip() if len(cells) > 5 else ""

            if title and "제목" not in title:
                results.append({
                    "기관": "정보통신산업진흥원(NIPA)", "공고구분": "입찰공고",
                    "제목": title, "담당자": author, "등록일": reg_date,
                    "조회수": views, "첨부파일": "", "원문링크": link,
                })
        if not results:
            Path("nipa_debug.html").write_text(page.content(), encoding="utf-8")
    except Exception as e:
        print(f"  [NIPA] 수집 실패: {e}")
    return results


# ============================================================
# 4. 행정안전부 (Playwright) — 실제로 절대경로 링크가 나옴을 확인, 안전장치만 추가
# ============================================================
def fetch_mois(page):
    results = []
    url = "https://www.mois.go.kr/frt/bbs/type013/commonSelectBoardList.do?bbsId=BBSMSTR_000000000006"
    try:
        page.goto(url, wait_until="networkidle", timeout=30000)
        page.wait_for_timeout(2000)
        try:
            page.wait_for_selector("table tbody tr", timeout=8000)
        except Exception:
            pass
        for row in page.query_selector_all("table tbody tr")[:15]:
            cells = row.query_selector_all("td")
            if len(cells) < 5:
                continue
            title_el = row.query_selector("a")
            title = title_el.inner_text().strip() if title_el else cells[1].inner_text().strip()
            link = title_el.get_attribute("href") if title_el else ""
            link = fix_link(link, "https://www.mois.go.kr")
            author = cells[-3].inner_text().strip() if len(cells) >= 3 else ""
            reg_date = cells[-2].inner_text().strip() if len(cells) >= 2 else ""
            views = cells[-1].inner_text().strip() if len(cells) >= 1 else ""
            has_attach = "1" if row.query_selector("img[alt*=첨부]") else ""
            if title:
                results.append({
                    "기관": "행정안전부", "공고구분": "공지/입찰", "제목": title,
                    "담당자": author, "등록일": reg_date, "조회수": views,
                    "첨부파일": "Y" if has_attach else "", "원문링크": link,
                })
        if not results:
            Path("mois_debug.html").write_text(page.content(), encoding="utf-8")
    except Exception as e:
        print(f"  [행정안전부] 수집 실패: {e}")
    return results


# ============================================================
# 5. AIHub — 목록 링크가 href="#" 뿐이라 클릭 캡처로 실제 URL 확인
# ============================================================
def fetch_aihub(page):
    results = []
    list_url = "https://www.aihub.or.kr/aihubnews/bsnspblanc/list.do?currMenu=133&topMenu=103"
    try:
        page.goto(list_url, wait_until="networkidle", timeout=30000)
        page.wait_for_timeout(2000)
        try:
            page.wait_for_selector("table tbody tr", timeout=8000)
        except Exception:
            pass
        row_count = len(page.query_selector_all("table tbody tr"))
        row_count = min(row_count, 10)

        for i in range(row_count):
            page.goto(list_url, wait_until="networkidle", timeout=30000)
            page.wait_for_timeout(1500)
            rows = page.query_selector_all("table tbody tr")
            if i >= len(rows):
                break
            row = rows[i]
            cells = row.query_selector_all("td")
            if len(cells) < 3:
                continue
            title_el = row.query_selector("a") or cells[0]
            title = title_el.inner_text().strip()
            reg_date = cells[-2].inner_text().strip() if len(cells) >= 2 else ""
            views = cells[-1].inner_text().strip() if len(cells) >= 1 else ""

            link = click_and_capture_url(page, title_el, list_url)

            if title:
                results.append({
                    "기관": "AI허브(AIHub)", "공고구분": "사업공고", "제목": title,
                    "등록일": reg_date, "조회수": views, "원문링크": link,
                })
        if not results:
            Path("aihub_debug.html").write_text(page.content(), encoding="utf-8")
    except Exception as e:
        print(f"  [AIHub] 수집 실패: {e}")
    return results


# ============================================================
# 6. 국가AI전략위원회 — 클라이언트 렌더링, 대기시간 연장 + 클릭 캡처
# ============================================================
def fetch_aikorea(page):
    results = []
    list_url = "https://www.aikorea.go.kr/web/board/brdList.do?menu_cd=000010"
    try:
        page.goto(list_url, wait_until="networkidle", timeout=30000)
        page.wait_for_timeout(4000)  # 클라이언트 렌더링 대기 연장
        try:
            page.wait_for_selector("table tbody tr", timeout=10000)
        except Exception:
            pass
        row_count = len(page.query_selector_all("table tbody tr"))
        row_count = min(row_count, 10)

        if row_count == 0:
            Path("aikorea_debug.html").write_text(page.content(), encoding="utf-8")
            print("  [국가AI전략위원회] 목록이 비어 있음 — aikorea_debug.html 확인 필요 "
                  "(클라이언트 렌더링 지연 또는 셀렉터 변경 가능성)")
            return results

        for i in range(row_count):
            page.goto(list_url, wait_until="networkidle", timeout=30000)
            page.wait_for_timeout(3000)
            rows = page.query_selector_all("table tbody tr")
            if i >= len(rows):
                break
            row = rows[i]
            cells = row.query_selector_all("td")
            if len(cells) < 2:
                continue
            title_el = row.query_selector("a") or cells[0]
            title = title_el.inner_text().strip()
            reg_date = cells[-2].inner_text().strip() if len(cells) >= 2 else ""
            views = cells[-1].inner_text().strip() if len(cells) >= 1 else ""

            link = click_and_capture_url(page, title_el, list_url)

            if title:
                results.append({
                    "기관": "국가인공지능전략위원회", "공고구분": "일반공지", "제목": title,
                    "등록일": reg_date, "조회수": views, "원문링크": link,
                })
    except Exception as e:
        print(f"  [국가AI전략위원회] 수집 실패: {e}")
    return results


# ============================================================
# 7. IITP
# ============================================================
def fetch_iitp(page):
    results = []
    url = ("https://www.iitp.kr/web/lay1/program/S1T44C51/iris/list.do"
           "?cms_menu_seq=51&cpage=1&rows=20&keyword=&condition=&sort=latest")
    try:
        page.goto(url, wait_until="networkidle", timeout=30000)
        page.wait_for_timeout(3000)
        try:
            page.wait_for_selector(".board_list_area li", state="attached", timeout=8000)
        except Exception:
            pass
        candidates = page.query_selector_all(".board_list_area li")
        if not candidates:
            candidates = page.query_selector_all("li")
        for item in candidates:
            title_el = item.query_selector("a")
            if not title_el:
                continue
            text = item.inner_text().strip()
            if "접수기간" not in text:
                continue
            title = title_el.inner_text().strip()
            link = title_el.get_attribute("href") or ""
            if link.startswith("./"):
                link = "https://www.iitp.kr/web/lay1/program/S1T44C51/iris/" + link[2:]
            elif link.startswith("/"):
                link = "https://www.iitp.kr" + link
            m = re.search(r"(\d{4}-\d{2}-\d{2})\s*~\s*(\d{4}-\d{2}-\d{2})", text)
            start, end = (m.group(1), m.group(2)) if m else ("", "")
            charger_m = re.search(r"담당자\s*([가-힣\*]+)", text)
            charger = charger_m.group(1) if charger_m else ""
            title_clean = re.sub(r"접수기간.*", "", title).strip() or title
            results.append({
                "기관": "정보통신기획평가원(IITP)", "공고구분": "입찰공고",
                "제목": title_clean, "담당자": charger, "등록일": start,
                "마감일": end, "원문링크": link,
            })
        if not results:
            Path("iitp_debug.html").write_text(page.content(), encoding="utf-8")
    except Exception as e:
        print(f"  [IITP] 수집 실패: {e}")
    return filter_active(results)


# ============================================================
# 8. NIA
# ============================================================
def fetch_nia(page):
    results = []
    url = "https://www.nia.or.kr/site/nia_kor/ex/bbs/List.do?cbIdx=78336"
    try:
        page.goto(url, wait_until="networkidle", timeout=30000)
        page.wait_for_timeout(2000)
        for a in page.query_selector_all("div.board_type01 li a[onclick]"):
            onclick = a.get_attribute("onclick") or ""
            m = re.search(r"doBbsFView\('(\d+)','(\d+)','(\d+)','(\d+)'\)", onclick)
            if not m:
                continue
            cb_idx, bc_idx, _, parent_seq = m.groups()
            title = re.sub(r"\(새 ?글\)", "", a.get_attribute("title") or "").strip()
            title = re.sub(r"-?첨부파일\s*있음", "", title).strip()
            em_texts = [e.inner_text().strip() for e in a.query_selector_all("em")]
            reg_date = next((t for t in em_texts if re.match(r"\d{4}\.\d{2}\.\d{2}", t)), "")
            link = f"https://www.nia.or.kr/site/nia_kor/ex/bbs/View.do?cbIdx={cb_idx}&bcIdx={bc_idx}&parentSeq={parent_seq}"
            if title:
                results.append({
                    "기관": "한국지능정보사회진흥원(NIA)", "공고구분": "입찰공고",
                    "제목": title, "등록일": reg_date,
                    "조회수": "", "첨부파일": "Y", "원문링크": link,
                })
        if not results:
            Path("nia_debug.html").write_text(page.content(), encoding="utf-8")
    except Exception as e:
        print(f"  [NIA] 수집 실패: {e}")
    return results


# ============================================================
# 9. KERIS — 목록 "번호" ≠ 상세 tenderSeq 이므로 클릭 캡처로 실제 URL 확인
#    담당자/담당부서는 목록 페이지 전체 공통값(재무회계부)이며 개별 공고 담당자가 아님
# ============================================================
def fetch_keris(page):
    results = []
    list_url = "https://www.keris.or.kr/main/tender/view/selectTenderList.do?mi=1076"
    try:
        page.goto(list_url, wait_until="networkidle", timeout=30000)
        page.wait_for_timeout(2000)
        try:
            page.wait_for_selector("table tbody tr", timeout=8000)
        except Exception:
            pass

        # 페이지 공통 담당자 정보 (모든 공고에 동일하게 표시되는 값)
        dept, manager = "", ""
        try:
            info_text = page.inner_text("body")
            m_dept = re.search(r"담당자\s*([가-힣]+부|[가-힣]+팀|[가-힣]+과)\s*([가-힣]{2,4})", info_text)
            if m_dept:
                dept, manager = m_dept.group(1), m_dept.group(2)
        except Exception:
            pass

        row_count = len(page.query_selector_all("table tbody tr"))
        row_count = min(row_count, 10)

        for i in range(row_count):
            page.goto(list_url, wait_until="networkidle", timeout=30000)
            page.wait_for_timeout(1500)
            rows = page.query_selector_all("table tbody tr")
            if i >= len(rows):
                break
            row = rows[i]
            cells = row.query_selector_all("td")
            if len(cells) < 4:
                continue
            title_raw = cells[1].inner_text().strip()
            title_clean = re.sub(r"새글", "", title_raw).strip()
            reg_date = cells[2].inner_text().strip()
            due_date = cells[3].inner_text().strip()

            clickable = cells[1].query_selector("a") or cells[1]
            link = click_and_capture_url(page, clickable, list_url)

            if title_clean:
                results.append({
                    "기관": "한국교육학술정보원(KERIS)", "공고구분": "입찰공고",
                    "제목": title_clean, "담당부서": dept, "담당자": manager,
                    "등록일": reg_date, "마감일": due_date, "원문링크": link,
                })
        if not results:
            Path("keris_debug.html").write_text(page.content(), encoding="utf-8")
    except Exception as e:
        print(f"  [KERIS] 수집 실패: {e}")
    return filter_active(results)


# ============================================================
# 공통 스타일/유틸
# ============================================================
def style_header(ws, columns):
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def auto_width(ws, columns):
    for i, col_name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = COL_WIDTH_HINTS.get(col_name, 20)


def find_col_index(columns, name):
    return columns.index(name) + 1 if name in columns else None


def apply_classification(item):
    """biz_classifier 결과를 item 딕셔너리에 채워넣는다."""
    result = classify_and_score(item.get("제목", ""))
    item["등급"] = result["등급"]
    item["카테고리"] = result["카테고리"]
    item["추천솔루션"] = result.get("추천솔루션", "-")
    item["매칭키워드"] = result["매칭키워드"]
    return item


def apply_row_style(ws, row_idx, columns):
    """등급(상/중/하) 기반 스타일링 + 마감 임박 3일 이내는 강조 테두리."""
    grade_idx = find_col_index(columns, "등급")
    due_idx = find_col_index(columns, "마감일")

    grade = ws.cell(row_idx, grade_idx).value if grade_idx else "하"
    due = ws.cell(row_idx, due_idx).value if due_idx else None

    deadline_date = parse_date(due) if due_idx else None
    is_urgent = False
    if deadline_date:
        days_left = (deadline_date - date.today()).days
        is_urgent = 0 <= days_left <= 3

    fill = GRADE_FILL.get(grade, GRADE_FILL["하"])
    font = GRADE_FONT.get(grade, GRADE_FONT["하"])

    for c in range(1, len(columns) + 1):
        cell = ws.cell(row_idx, c)
        cell.fill = fill
        cell.font = font
        cell.border = URGENT_BORDER if is_urgent else NO_BORDER

    return grade == "상", is_urgent


def recompute_highlight_counts(ws, columns):
    highlight = urgent = 0
    title_idx = find_col_index(columns, "제목")
    if not title_idx:
        return 0, 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, title_idx).value is None:
            continue
        is_p, is_u = apply_row_style(ws, r, columns)
        highlight += 1 if is_p else 0
        urgent += 1 if is_u else 0
    return highlight, urgent


def update_row_if_needed(ws, row_idx, item, columns):
    changed = False
    due_idx = find_col_index(columns, "마감일")
    title_idx = find_col_index(columns, "제목")
    link_idx = find_col_index(columns, "원문링크")
    gubun_idx = find_col_index(columns, "공고구분")

    if due_idx:
        old_due = ws.cell(row_idx, due_idx).value or ""
        new_due = item.get("마감일", "") or ""
        if new_due and new_due != old_due:
            ws.cell(row_idx, due_idx, new_due)
            changed = True

    if title_idx:
        old_title = ws.cell(row_idx, title_idx).value or ""
        new_title = item.get("제목", "") or ""
        if new_title and len(new_title) > len(old_title):
            ws.cell(row_idx, title_idx, new_title)
            changed = True

    if link_idx:
        old_link = ws.cell(row_idx, link_idx).value or ""
        new_link = item.get("원문링크", "") or ""
        if new_link and new_link != old_link:
            ws.cell(row_idx, link_idx, new_link)
            changed = True

    # 분류 결과도 최신으로 갱신
    for col in CLASSIFY_COLUMNS:
        col_idx = find_col_index(columns, col)
        if col_idx and item.get(col) is not None:
            ws.cell(row_idx, col_idx, item.get(col))

    if changed and gubun_idx:
        old_gubun = ws.cell(row_idx, gubun_idx).value or ""
        if "(갱신)" not in old_gubun:
            ws.cell(row_idx, gubun_idx, f"{old_gubun}(갱신)")
    return changed


def upsert_sheet(wb, sheet_name, items, columns):
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        existing_header = [c.value for c in ws[1]] if ws.max_row >= 1 else []
        if existing_header != columns:
            wb.remove(ws)
            ws = wb.create_sheet(sheet_name)
    else:
        ws = wb.create_sheet(sheet_name)

    if ws.max_row == 1 and all(c.value is None for c in ws[1]):
        for idx, col in enumerate(columns, start=1):
            ws.cell(row=1, column=idx, value=col)
    style_header(ws, columns)

    title_idx = find_col_index(columns, "제목")
    agency_idx = find_col_index(columns, "기관")
    existing_by_key = {}
    for r in range(2, ws.max_row + 1):
        agency = ws.cell(r, agency_idx).value if agency_idx else None
        title = ws.cell(r, title_idx).value if title_idx else None
        if agency and title:
            existing_by_key[(agency, normalize_title(title))] = r

    new_count = updated_count = 0

    for item in reversed(items):
        if not item.get("제목"):
            continue
        # ── biz_classifier 적용 ──
        apply_classification(item)

        key = (item.get("기관", ""), normalize_title(item["제목"]))
        if key in existing_by_key:
            row_idx = existing_by_key[key]
            if update_row_if_needed(ws, row_idx, item, columns):
                updated_count += 1
            continue

        ws.insert_rows(2)
        existing_by_key = {k: v + 1 for k, v in existing_by_key.items()}
        for col_idx, col_name in enumerate(columns, start=1):
            ws.cell(row=2, column=col_idx, value=item.get(col_name, ""))
        existing_by_key[key] = 2
        new_count += 1

    auto_width(ws, columns)
    return {"new": new_count, "updated": updated_count, "highlight": 0, "urgent": 0}


def archive_and_cleanup(wb):
    if ARCHIVE_SHEET in wb.sheetnames:
        arch_ws = wb[ARCHIVE_SHEET]
    else:
        arch_ws = wb.create_sheet(ARCHIVE_SHEET)

    if arch_ws.max_row == 1 and all(c.value is None for c in arch_ws[1]):
        for idx, col in enumerate(ARCHIVE_COLUMNS_BASE, start=1):
            arch_ws.cell(row=1, column=idx, value=col)
    style_header(arch_ws, ARCHIVE_COLUMNS_BASE)

    today = date.today()
    archived_count = 0

    for sheet_name in list(wb.sheetnames):
        if sheet_name in (ARCHIVE_SHEET, DASHBOARD_SHEET):
            continue
        if sheet_name not in SHEET_COLUMNS:
            continue
        columns = SHEET_COLUMNS[sheet_name]
        ws = wb[sheet_name]
        due_idx = find_col_index(columns, "마감일")
        if not due_idx:
            continue
        agency_idx = find_col_index(columns, "기관")
        gubun_idx = find_col_index(columns, "공고구분")
        title_idx = find_col_index(columns, "제목")
        reg_idx = find_col_index(columns, "등록일")
        link_idx = find_col_index(columns, "원문링크")

        rows_to_delete = []
        for r in range(2, ws.max_row + 1):
            due_val = ws.cell(r, due_idx).value
            due_date = parse_date(due_val)
            if due_date and (today - due_date).days > 7:
                arch_ws.append([
                    ws.cell(r, agency_idx).value if agency_idx else "",
                    ws.cell(r, gubun_idx).value if gubun_idx else "",
                    ws.cell(r, title_idx).value if title_idx else "",
                    ws.cell(r, reg_idx).value if reg_idx else "",
                    due_val,
                    ws.cell(r, link_idx).value if link_idx else "",
                    sheet_name,
                    today.strftime("%Y-%m-%d"),
                ])
                for c in range(1, len(ARCHIVE_COLUMNS_BASE) + 1):
                    cell = arch_ws.cell(arch_ws.max_row, c)
                    cell.fill = NO_FILL
                    cell.font = Font()
                    cell.border = NO_BORDER
                rows_to_delete.append(r)
                archived_count += 1
        for r in sorted(rows_to_delete, reverse=True):
            ws.delete_rows(r)

    delete_rows = []
    archived_col = len(ARCHIVE_COLUMNS_BASE)
    for r in range(2, arch_ws.max_row + 1):
        archived_on = arch_ws.cell(r, archived_col).value
        archived_date = parse_date(archived_on)
        if archived_date and (today - archived_date).days > 30:
            delete_rows.append(r)
    for r in sorted(delete_rows, reverse=True):
        arch_ws.delete_rows(r)

    auto_width(arch_ws, ARCHIVE_COLUMNS_BASE)
    return archived_count, len(delete_rows)


def build_dashboard(wb, stats, archived_count, deleted_count):
    if DASHBOARD_SHEET in wb.sheetnames:
        wb.remove(wb[DASHBOARD_SHEET])
    ws = wb.create_sheet(DASHBOARD_SHEET, 0)

    headers = ["기관 시트", "신규 추가", "갱신(재공고 등)", "상위등급(제품연관)", "마감 임박(3일 이내)", "최종 수집 시각"]
    for idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=idx, value=h)
    style_header(ws, headers)

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    row = 2
    totals = {"new": 0, "updated": 0, "highlight": 0, "urgent": 0}
    for sheet_name, s in stats.items():
        ws.cell(row, 1, sheet_name)
        ws.cell(row, 2, s.get("new", 0))
        ws.cell(row, 3, s.get("updated", 0))
        ws.cell(row, 4, s.get("highlight", 0))
        ws.cell(row, 5, s.get("urgent", 0))
        ws.cell(row, 6, now_str)
        for k in totals:
            totals[k] += s.get(k, 0)
        row += 1

    for c, val in zip(range(1, 6), ["합계", totals["new"], totals["updated"], totals["highlight"], totals["urgent"]]):
        cell = ws.cell(row, c, val)
        cell.font = Font(bold=True)
    row += 2
    ws.cell(row, 1, f"이번 실행에서 보관 시트로 이동된 공고: {archived_count}건")
    row += 1
    ws.cell(row, 1, f"보관 30일 경과로 영구 삭제된 공고: {deleted_count}건")

    for i, w in enumerate([24, 12, 16, 18, 16, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    print(f"=== Gov-Tracker 통합 수집 시작 ({datetime.now().strftime('%H:%M:%S')}) ===\n")
    all_by_sheet = {}

    print("[조달청] 수집 중...")
    all_by_sheet["조달청"] = fetch_g2b()
    print(f"  -> {len(all_by_sheet['조달청'])}건 수집")

    print("[알리오] 수집 중...")
    all_by_sheet["알리오"] = fetch_alio()
    print(f"  -> {len(all_by_sheet['알리오'])}건 수집")

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page()
        page.set_default_timeout(30000)

        print("[행정안전부] 수집 중...")
        all_by_sheet["행정안전부"] = fetch_mois(page)
        print(f"  -> {len(all_by_sheet['행정안전부'])}건 수집")

        print("[NIPA] 수집 중...")
        all_by_sheet["NIPA"] = fetch_nipa(page)
        print(f"  -> {len(all_by_sheet['NIPA'])}건 수집")

        print("[AIHub] 수집 중... (클릭 캡처 방식, 시간이 조금 걸립니다)")
        all_by_sheet["AIHub"] = fetch_aihub(page)
        print(f"  -> {len(all_by_sheet['AIHub'])}건 수집")

        print("[국가인공지능전략위원회] 수집 중... (클릭 캡처 방식)")
        all_by_sheet["국가AI전략위원회"] = fetch_aikorea(page)
        print(f"  -> {len(all_by_sheet['국가AI전략위원회'])}건 수집")

        print("[IITP] 수집 중...")
        all_by_sheet["IITP"] = fetch_iitp(page)
        print(f"  -> {len(all_by_sheet['IITP'])}건 수집")

        print("[NIA] 수집 중...")
        all_by_sheet["NIA"] = fetch_nia(page)
        print(f"  -> {len(all_by_sheet['NIA'])}건 수집")

        print("[KERIS] 수집 중... (클릭 캡처 방식)")
        all_by_sheet["KERIS"] = fetch_keris(page)
        print(f"  -> {len(all_by_sheet['KERIS'])}건 수집")

        browser.close()

    print("\n=== 엑셀 저장 중 ===")
    path = Path(EXCEL_FILE)
    wb = openpyxl.load_workbook(path) if path.exists() else openpyxl.Workbook()
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
        del wb["Sheet"]

    stats = {}
    for sheet_name, items in all_by_sheet.items():
        columns = SHEET_COLUMNS[sheet_name]
        s = upsert_sheet(wb, sheet_name, items, columns)
        stats[sheet_name] = s
        print(f"  [{sheet_name}] 신규 {s['new']}건, 갱신 {s['updated']}건 (수집 {len(items)}건 중)")

    archived_count, deleted_count = archive_and_cleanup(wb)
    print(f"  [보관 처리] 이동 {archived_count}건, 30일경과 영구삭제 {deleted_count}건")

    for sheet_name in stats.keys():
        if sheet_name in wb.sheetnames:
            hl, urg = recompute_highlight_counts(wb[sheet_name], SHEET_COLUMNS[sheet_name])
            stats[sheet_name]["highlight"] = hl
            stats[sheet_name]["urgent"] = urg

    build_dashboard(wb, stats, archived_count, deleted_count)

    wb.save(EXCEL_FILE)
    total_new = sum(s["new"] for s in stats.values())
    print(f"\n=== 완료: '{EXCEL_FILE}'에 총 {total_new}건의 신규 공고 추가 (요약 시트 확인) ===")


if __name__ == "__main__":
    main()
