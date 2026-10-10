# integrated_export.py
# 업체 중심 통합 엑셀: 낙찰 + 계약 + 쇼핑몰 납품(솔루션) 을 한 표로
#   · 업체검색: 업체명 일부 입력 → 그 업체가 진행 중·진행했던 사업/납품 전부
#   · 경쟁사 수주 현황 · API 관련 사업 수주 업체 · 대학 사업 현황(국립/사립) · 영업 파트너 후보
# 실행: python integrated_export.py --out integrated.xlsx [--dlvr-cache raw_cache.json]
import argparse
import io
import json
import re
import sys
from datetime import datetime

import pandas as pd
from sqlalchemy import text

import common  # noqa: F401  (한국시간 고정)
from common import solution_hits, is_competitor_match, detect_regions, competitor_product_maker
import contract_export as ce
import store
from db2 import get_engine

MY_REGIONS = {"서울", "인천", "강원", "전북", "전남", "광주", "제주"}
API_RE = re.compile(r"(API|에이피아이|게이트웨이|gateway|오픈\s*API|MSA|마이크로서비스|트래픽\s*(제어|관리)|대기열|유량\s*제어)", re.I)
UNIV_RE = re.compile(r"(대학교|대학원|산학협력단|교육대학|한국과학기술원|광주과학기술원|대구경북과학기술원|울산과학기술원|KAIST|GIST|DGIST|UNIST|폴리텍|대학\b|대학$)")
NATIONAL_UNIV = [
    "서울대학교", "부산대학교", "경북대학교", "전남대학교", "전북대학교", "충남대학교", "충북대학교", "강원대학교", "경상국립대학교",
    "제주대학교", "부경대학교", "한국해양대학교", "목포해양대학교", "목포대학교", "순천대학교", "군산대학교", "안동대학교",
    "창원대학교", "공주대학교", "금오공과대학교", "한밭대학교", "한국교통대학교", "강릉원주대학교", "한경국립대학교", "한경대학교",
    "한국교원대학교", "한국체육대학교", "서울과학기술대학교", "한국방송통신대학교", "인천대학교", "서울시립대학교",
    "교육대학교", "한국과학기술원", "광주과학기술원", "대구경북과학기술원", "울산과학기술원", "KAIST", "GIST", "DGIST", "UNIST", "한국전통문화대학교", "한국예술종합학교", "국립",
    "한국폴리텍", "한국기술교육대학교", "한국농수산대학교", "경찰대학", "육군사관학교", "해군사관학교", "공군사관학교", "국군간호사관학교",
]
COMMON_COLS = ["출처", "구분", "상태", "남은일수", "사업명", "제품", "사업명·제품", "수요기관", "기관구분", "대학", "지역", "담당지역", "업체명", "대표업체",
               "금액", "일자", "종료일", "종료추정", "자사관련", "자사관련단어", "경쟁사", "API관련", "원문", "검색용업체명", "_gk", "_key"]


SIDO_SHORT = {"서울특별시": "서울", "서울": "서울", "부산광역시": "부산", "부산": "부산", "대구광역시": "대구", "대구": "대구",
              "인천광역시": "인천", "인천": "인천", "광주광역시": "광주", "광주": "광주", "대전광역시": "대전", "대전": "대전",
              "울산광역시": "울산", "울산": "울산", "세종특별자치시": "세종", "세종": "세종", "경기도": "경기", "경기": "경기",
              "강원특별자치도": "강원", "강원도": "강원", "강원": "강원", "충청북도": "충북", "충북": "충북", "충청남도": "충남", "충남": "충남",
              "전북특별자치도": "전북", "전라북도": "전북", "전북": "전북", "전라남도": "전남", "전남": "전남",
              "경상북도": "경북", "경북": "경북", "경상남도": "경남", "경남": "경남", "제주특별자치도": "제주", "제주": "제주"}
CENTRAL_RE = re.compile(r"(부|처|청|위원회|감사원|국회|대법원|헌법재판소|법원행정처|대통령)$|^(기획재정부|교육부|과학기술정보통신부|외교부|통일부|법무부|국방부|행정안전부|문화체육관광부|농림축산식품부|산업통상자원부|보건복지부|환경부|고용노동부|여성가족부|국토교통부|해양수산부|중소벤처기업부)")


def _univ_index():
    """대학알리미 학교개황(data/univ_master.csv): 학교명 → (지역, 설립구분)"""
    import csv
    idx = {}
    try:
        for r in csv.DictReader(open("data/univ_master.csv", encoding="utf-8-sig")):
            nm = re.sub(r"\s+", "", r["학교명"])
            idx.setdefault(nm, (r["지역"], r["설립구분"]))
    except Exception:
        pass
    return dict(sorted(idx.items(), key=lambda x: -len(x[0])))


UNIV_IDX = _univ_index()


def univ_lookup(name):
    n = re.sub(r"\s+", "", str(name or ""))
    for k, v in UNIV_IDX.items():
        if k in n:
            return v
    return None


# 자주 나오는 공공기관 본사 소재지 (조달청 수요기관 지역표가 채워지기 전 임시 보완)
KNOWN_HQ = {"한국지능정보사회진흥원": "대구", "한국항공우주연구원": "대전", "한국건설기술연구원": "경기", "한국교육학술정보원": "대구",
            "한국과학기술정보연구원": "대전", "한국인터넷진흥원": "전남", "국방과학연구소": "대전", "국립보건연구원": "충북",
            "한국생산기술연구원": "충남", "한국지역정보개발원": "서울", "중소벤처기업진흥공단": "경남", "한국고용정보원": "충북",
            "한국교통안전공단": "경북", "각 수요기관": "중앙(전국)"}
CODE_REGION = {}


def load_code_region():
    try:
        with get_engine().begin() as conn:
            rows = conn.execute(text("SELECT code, region FROM dminstt_region WHERE region <> ''")).fetchall()
        CODE_REGION.update({str(c): r for c, r in rows})
    except Exception:
        pass
    print(f"[지역표] 수요기관 코드 {len(CODE_REGION):,}곳")


def resolve_region(name, hint="", code=""):
    """기관 소재 시·도 (짧은 이름). 순서: ① 원자료 지역 ② 대학 목록 ③ 기관명 속 지역명 ④ 중앙부처 ⑤ 미확인"""
    if code and str(code) in CODE_REGION:
        return CODE_REGION[str(code)]
    for k, v in KNOWN_HQ.items():
        if k in str(name or ""):
            return v
    h = str(hint or "").strip()
    if h:
        tok = h.split()[0]
        if tok in SIDO_SHORT:
            return SIDO_SHORT[tok]
    u = univ_lookup(name)
    if u and u[0]:
        return SIDO_SHORT.get(u[0], u[0])
    r = detect_regions(name)
    if r:
        return r[0]
    if CENTRAL_RE.search(str(name or "").strip()):
        return "중앙(전국)"
    return "미확인"


UNIV_EXCLUDE = re.compile(r"(병원|협의회|협회|진흥원|연구재단|장학재단|해양과학기술원|부설)")


def univ_type(name):
    n = str(name or "")
    if not UNIV_RE.search(n) or UNIV_EXCLUDE.search(n):
        return ""
    u = univ_lookup(n)
    if u:
        return "사립" if u[1] == "사립" else "국립·공립"
    return "국립·공립" if any(k in n for k in NATIONAL_UNIV) else "사립"


def _status(end, today):
    if pd.isna(end):
        return "종료일 미상", None
    d = (end - today).days
    return ("완료" if d < 0 else "곧 완료" if d <= ce.SOON_DAYS else "진행중"), d


def _common(df, comp):
    """공통 칼럼 채우기 (각 출처에서 사업명·수요기관·업체목록·금액·일자·종료일을 만든 뒤 호출)"""
    today = pd.Timestamp(datetime.now().date())
    if "제품" not in df:
        df["제품"] = ""
    df["사업명·제품"] = (df["사업명"].astype(str) + " " + df["제품"].astype(str)).str.strip()
    end = pd.to_datetime(df["종료일"], errors="coerce")
    st = [_status(e, today) for e in end]
    df["상태"] = [s for s, _ in st]
    df["남은일수"] = [d for _, d in st]
    df["업체명"] = df["업체목록"].map(lambda xs: " / ".join(xs))
    df["대표업체"] = df["업체목록"].map(lambda xs: xs[0] if xs else "")
    df["검색용업체명"] = df["업체목록"].map(ce.search_text)
    df["_gk"] = df["대표업체"].map(ce.group_key)
    df["대학"] = df["수요기관"].map(univ_type)
    hints = df["_지역힌트"] if "_지역힌트" in df else pd.Series([""] * len(df), index=df.index)
    codes = df["_기관코드"] if "_기관코드" in df else pd.Series([""] * len(df), index=df.index)
    df["지역"] = [resolve_region(n, h, c) for n, h, c in zip(df["수요기관"], hints, codes)]
    df["담당지역"] = df["지역"].map(lambda r: "Y" if r in MY_REGIONS else "")
    df["기관구분"] = df["대학"].map(lambda u: "대학" if u else "공공")
    hits = df["사업명·제품"].map(solution_hits)
    df["자사관련단어"] = hits.map(lambda h: ", ".join(dict.fromkeys(h)) if h else "")
    df["자사관련"] = (df["자사관련단어"] != "").map({True: "Y", False: ""})
    def _comp_name(company, text_):
        for nm in comp:
            if is_competitor_match(company, [nm]):
                return nm
        return competitor_product_maker(text_)
    df["경쟁사명"] = [_comp_name(c_, t_) for c_, t_ in zip(df["업체명"], df["사업명·제품"])]
    df["경쟁사"] = (df["경쟁사명"] != "").map({True: "Y", False: ""})
    df["API관련"] = df["사업명·제품"].map(lambda t: "Y" if API_RE.search(str(t)) else "")
    return df


def load_contracts(comp):
    raw = ce.load_rows()
    if raw.empty:
        return pd.DataFrame(columns=COMMON_COLS), set()
    d = ce.prepare(raw, competitors=comp)
    out = pd.DataFrame({
        "출처": "계약", "구분": d["biz_type"], "사업명": d["title"], "제품": "", "수요기관": d["수요기관"],
        "업체목록": d["업체목록"], "금액": d["계약금액"], "일자": d["cntrct_date"], "종료일": d["end_date"],
        "종료추정": d["end_est"], "원문": d["url"], "_key": "계약|" + d["uniq_key"],
        "_기관코드": d["dminstt_raw"].map(lambda x: (re.findall(r"\[\d+\^([^\^\]]+)\^", str(x)) or [""])[0]),
    })
    ntce = {str(x).split("-")[0] for x in d["ntce_no"] if str(x)}
    return _common(out, comp), ntce


def load_awards(comp, skip_bid_nos):
    try:
        with get_engine().begin() as conn:
            a = pd.read_sql(text("SELECT * FROM g2b_awards"), conn).fillna("")
    except Exception:
        return pd.DataFrame(columns=COMMON_COLS)
    if a.empty:
        return pd.DataFrame(columns=COMMON_COLS)
    # 같은 공고가 계약으로도 들어와 있으면 계약 쪽을 씀 (계약기간이 정확함)
    a = a[~a["bid_no"].isin(skip_bid_nos)]
    a = a.sort_values("updated_at", ascending=False).drop_duplicates(subset=["bid_no", "bid_ord", "winner"], keep="first")
    out = pd.DataFrame({
        "출처": "낙찰", "구분": a["biz_type"], "사업명": a["title"], "제품": "",
        "수요기관": [d or n for d, n in zip(a["dminstt"], a["ntce_instt"])],
        "업체목록": a["winner"].map(lambda w: [w] if w else []), "금액": a["amount"].map(ce._won),
        "일자": a["event_date"], "종료일": a["end_date"], "종료추정": "Y", "원문": a["url"], "_key": "낙찰|" + a["uniq_key"],
        "_기관코드": a["dminstt_cd"],
    })
    return _common(out, comp)


def load_dlvr(comp, cache_path):
    import dlvr_export as de
    import os
    if not os.path.exists(cache_path):
        return pd.DataFrame(columns=COMMON_COLS)
    de.OFFLINE = True
    de.CACHE_FILE = cache_path
    rows = de.collect()
    if not rows:
        return pd.DataFrame(columns=COMMON_COLS)
    d = pd.DataFrame(rows).fillna("")
    d["_chg"] = pd.to_numeric(d.get("dlvrReqChgOrd", 0), errors="coerce").fillna(0)
    d = d.sort_values("_chg", ascending=False).drop_duplicates(subset=["dlvrReqNo", "prdctSno", "prdctIdntNo"], keep="first")
    spec = d["prdctIdntNoNm"].astype(str)
    out = pd.DataFrame({
        "출처": "쇼핑몰 납품", "구분": "쇼핑몰(" + d["_구분"].astype(str) + ")",
        "사업명": d["dlvrReqNm"].astype(str),
        "제품": spec.map(lambda s: " · ".join([x.strip() for x in s.split(",")[1:4]]) or s),
        "수요기관": d["dminsttNm"], "업체목록": d["corpNm"].map(lambda w: [w] if w else []), "금액": d["prdctAmt"].map(ce._won),
        "일자": d["dlvrReqRcptDate"].map(ce._date),
        # 솔루션 납품은 '납품일 + 1년'을 유지보수·재구매 시점으로 봄 (추정)
        "종료일": (pd.to_datetime(d["dlvrReqRcptDate"].map(ce._date), errors="coerce") + pd.Timedelta(days=365)).dt.strftime("%Y-%m-%d"),
        "종료추정": "Y", "원문": "", "_key": "납품|" + d["dlvrReqNo"].astype(str) + "|" + d["prdctSno"].astype(str),
        "_지역힌트": d["dminsttRgnNm"], "_기관코드": d["dminsttCd"],
    })
    out = _common(out, comp)
    # 경쟁사·비교군 '제품' 납품은 판매업체가 리셀러여도 경쟁사 수주로 표시
    kind = d["_구분"].tolist()
    out["경쟁사명"] = [n_ or (competitor_product_maker(sp) if k == "경쟁사" else "") for n_, k, sp in zip(out["경쟁사명"], kind, d["prdctIdntNoNm"])]
    out["경쟁사"] = (out["경쟁사명"] != "").map({True: "Y", False: ""})
    out["API관련"] = ["Y" if (a == "Y" or "API" in str(f)) else "" for a, f in zip(out["API관련"], d["_제품군"])]
    out.loc[[k == "자사" for k in kind], "자사관련"] = "Y"
    return out


def load_univ_bids(comp):
    """사립대·전문대 홈페이지 입찰 게시판에서 모은 글 (univ_crawl.py)"""
    try:
        with get_engine().begin() as conn:
            u = pd.read_sql(text("SELECT * FROM univ_bids"), conn).fillna("")
    except Exception:
        return pd.DataFrame(columns=COMMON_COLS)
    if u.empty:
        return pd.DataFrame(columns=COMMON_COLS)
    ev = pd.to_datetime(u["date"], errors="coerce")
    for c in ("budget", "deadline", "period_end"):          # 예전 수집분에는 없는 칸
        if c not in u.columns:
            u[c] = ""
    corp_like = re.compile(r"(주식회사|\(주\)|㈜|유한|회사|시스템|정보|기술|테크|텍|소프트|솔루션|네트웍|네트워크|컴퍼니|커뮤니케이션|아이티|IT|디지털|데이타|데이터|산업|전자|통신|엔지니어링|corp|inc)", re.I)
    u["winner"] = u["winner"].map(lambda w: w if corp_like.search(str(w)) else "")
    # 금액: 낙찰·계약금액 → 없으면 공고 예산(추정가격·기초금액)
    amt = u["amount"].where(u["amount"].astype(str).str.strip() != "", u["budget"])
    # 사업 종료일: 본문·첨부의 계약(사업)기간 → 없으면 공고(결과)일 + 1년 추정
    pe = pd.to_datetime(u["period_end"], errors="coerce")
    est = (ev + pd.Timedelta(days=365)).dt.strftime("%Y-%m-%d")
    end = pe.dt.strftime("%Y-%m-%d").where(pe.notna(), est)
    kind = u["kind"].map(lambda k: {"공고": "입찰공고", "결과": "입찰결과", "계약공개": "계약공개"}.get(k, "입찰" + str(k)))
    out = pd.DataFrame({
        "출처": "대학 홈페이지", "구분": kind, "사업명": u["title"], "제품": "",
        "수요기관": u["school"] + u["campus"].map(lambda c: "" if c in ("", "본교") else f"({c})"),
        "업체목록": u["winner"].map(lambda w: [w] if w else []), "금액": amt.map(ce._won),
        "일자": u["date"], "종료일": end, "종료추정": pe.isna().map(lambda x: "Y" if x else ""),
        "원문": u["url"], "_key": "대학|" + u["uniq_key"], "_지역힌트": u["region"],
    })
    return _common(out, comp)


FIELD_RE = [  # 분야 판별 (위에서부터 먼저 맞는 것)
    ("예약·접속·대기열·매크로 (자사 연관)", re.compile(r"(예약|대기열|접속|매크로|수강신청|트래픽|티켓|선착순|부하\s*테스트|봇)")),
    ("AI·데이터", re.compile(r"(AI|인공지능|데이터|빅데이터|챗봇|LLM|생성형)", re.I)),
    ("정보보호", re.compile(r"(보안|정보보호|관제|개인정보|해킹|백신|침해)")),
    ("클라우드·인프라", re.compile(r"(클라우드|서버|네트워크|스토리지|인프라|이중화|DaaS|전산\s*장비|UPS)", re.I)),
    ("홈페이지·플랫폼·앱", re.compile(r"(홈페이지|누리집|포털|플랫폼|웹|앱|모바일)")),
    ("정보시스템 구축·고도화", re.compile(r"(구축|고도화|개발|차세대|ISP|ISMP|전환)")),
    ("유지관리·운영", re.compile(r"(유지관리|유지보수|운영|위탁)")),
]


def field_of(title):
    t = str(title or "")
    for name, rx in FIELD_RE:
        if rx.search(t):
            return name
    return "기타"


def build(all_df):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    def clean(v):
        return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v
    d = all_df.copy()
    for col in d.columns:
        if not pd.api.types.is_numeric_dtype(d[col]):
            d[col] = d[col].map(clean)

    # ---- 공통 칼럼 (모든 시트 같은 순서) ----
    today = pd.Timestamp(datetime.now().date())
    end = pd.to_datetime(d["종료일"], errors="coerce")
    days = (end - today).dt.days
    d["사업종료일"] = end.dt.strftime("%Y-%m-%d").fillna("")
    d["종료"] = "미상"
    d.loc[days < 0, "종료"] = "종료"
    d.loc[(days >= 0) & (days <= ce.SOON_DAYS), "종료"] = "종료 임박"
    d.loc[days > ce.SOON_DAYS, "종료"] = "진행중"
    d["종료 D-day"] = [("" if pd.isna(x) else (f"D-{int(x)}" if x > 0 else ("D-DAY" if x == 0 else f"D+{int(-x)}"))) for x in days]
    d["_days"] = days
    d["기관특성"] = d["대학"].map(lambda u: "대학" if u else "공공")
    d["대학구분"] = d["대학"].map(lambda u: {"국립·공립": "국·공립", "사립": "사립"}.get(u, ""))
    d["분야"] = d["사업명"].map(field_of)
    d["_month"] = pd.to_datetime(d["일자"], errors="coerce").dt.strftime("%Y-%m")
    d["솔루션연관"] = ((d["자사관련"] == "Y") | (d["API관련"] == "Y") | (d["경쟁사"] == "Y")
                    | (d["분야"].str.startswith("예약")) | (d["출처"] == "쇼핑몰 납품")).map({True: "Y", False: ""})
    order = {"종료 임박": 0, "진행중": 1, "미상": 2, "종료": 3}
    d = d.assign(_o=d["종료"].map(order), _dd=d["_days"].fillna(10 ** 6))
    d = d.sort_values(["_o", "_dd"]).reset_index(drop=True)

    wb = Workbook(write_only=True)
    navy = PatternFill("solid", fgColor="1B2A4A")
    hf = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)
    big = Font(bold=True, size=16)
    link_font = Font(color="1F4FD1", underline="single")
    yellow = PatternFill("solid", fgColor="FFF59D")
    soon_fill = PatternFill("solid", fgColor="FDECEA")
    EOK, WON = '#,##0.0"억"', '#,##0'
    BAR = "2F5BD3"

    def c(ws, v, font=None, fill=None, fmt=None, wrap=False):
        x = WriteOnlyCell(ws, value=v)
        if font:
            x.font = font
        if fill:
            x.fill = fill
        if fmt:
            x.number_format = fmt
        if wrap:
            x.alignment = Alignment(wrap_text=True, vertical="top")
        return x

    def _labels():
        lb = DataLabelList()
        lb.showVal, lb.showCatName, lb.showSerName, lb.showLegendKey, lb.showPercent = True, False, False, False, False
        return lb

    def money(ws, v):
        if v is None or v == "" or (isinstance(v, float) and pd.isna(v)):
            return c(ws, None)
        v = float(v)
        return c(ws, round(v / 1e8, 2), fmt=EOK) if abs(v) >= 1e7 else c(ws, int(v), fmt=WON)

    def link_cell(ws, url, title):
        t = clean(str(title or "")).replace('"', '""')[:240]
        u = clean(str(url or "")).replace('"', "%22")
        if u.startswith("http") and len(u) < 250:
            return c(ws, f'=HYPERLINK("{u}","{t}")', font=link_font)
        return c(ws, t)

    # 고정 5칸 + 사업 6칸
    VIEW = [("기관특성", "기관특성", 8), ("대학구분", "대학구분", 8), ("업체명", "업체명", 26), ("수요기관", "수요기관", 26),
            ("지역", "지역", 9), ("사업명 (누르면 원문)", "사업명", 48), ("제품", "제품", 28), ("금액", "금액", 11),
            ("사업종료일", "사업종료일", 11), ("종료", "종료", 9), ("종료 D-day", "종료 D-day", 9)]
    HELP = [("금액(원)", "_won", 14), ("원문 주소", "원문", 10), ("검색용 업체명", "검색용업체명", 10)]
    L = {t: get_column_letter(i + 1) for i, (t, _, _) in enumerate(VIEW + HELP)}

    def write_rows(ws, sub, helper=False, start_row=2):
        for rn, r in enumerate(sub.to_dict("records"), start_row):
            row = []
            hot = r["종료"] == "종료 임박"
            for t, k, _ in VIEW:
                v = r.get(k, "")
                if k == "사업명":
                    row.append(link_cell(ws, r.get("원문", ""), v))
                elif k == "금액":
                    row.append(money(ws, v))
                elif k in ("종료", "종료 D-day") and hot:
                    row.append(c(ws, v, fill=soon_fill))
                else:
                    row.append(v)
            if helper:
                w = r.get("금액")
                row += [c(ws, None if w is None or w == "" or pd.isna(w) else int(w), fmt=WON), r.get("원문", ""),
                        r.get("검색용업체명", ""),
                        f'=IF(업체검색!$F$3="","",IF(ISNUMBER(SEARCH(업체검색!$F$3,{L["검색용 업체명"]}{rn})),ROW(),""))']
            ws.append(row)

    def table(name, sub, helper=False, note=""):
        ws = wb.create_sheet(name)
        cols = VIEW + (HELP + [("검색키", "", 6)] if helper else [])
        for i, (_, _, w) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for t in ("금액(원)", "원문 주소", "검색용 업체명", "검색키") if helper else ():
            ws.column_dimensions[L.get(t, get_column_letter(len(VIEW) + len(HELP) + 1))].hidden = True
        top = 1
        if note:
            ws.append([c(ws, note, bold)])
            top = 2
        ws.append([c(ws, t, hf, navy) for t, _, _ in cols])
        ws.freeze_panes = f"F{top + 1}"
        write_rows(ws, sub, helper, start_row=top + 1)
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(VIEW))}{max(len(sub) + top, top + 1)}"
        return ws

    def grouped_table(name, sub, key, order_keys, note):
        """key별로 묶어 쓰고 각 묶음의 시작 행을 돌려줌 (대시보드 링크용)"""
        ws = wb.create_sheet(name)
        for i, (_, _, w) in enumerate(VIEW, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.append([c(ws, note, bold)])
        ws.append([c(ws, t, hf, navy) for t, _, _ in VIEW])
        ws.freeze_panes = "F3"
        anchors, rn = {}, 3
        for k in order_keys:
            part = sub[sub[key] == k]
            if part.empty:
                continue
            ws.append([c(ws, f"■ {k} — {len(part):,}건", bold, PatternFill("solid", fgColor="E8EEF9"))])
            anchors[k] = rn
            rn += 1
            write_rows(ws, part, start_row=rn)
            rn += len(part)
        return anchors

    # ---- 집계 ----
    sol = d[d["솔루션연관"] == "Y"]
    own_gk = ce.group_key("에스티씨랩")
    top_comp = sol[(sol["대표업체"] != "") & (sol["_gk"] != own_gk)].groupby("_gk").agg(
        업체명=("대표업체", lambda x: x.value_counts().index[0]), 건수=("_key", "size"),
        진행=("종료", lambda x: int(x.isin(["진행중", "종료 임박"]).sum()))).sort_values("건수", ascending=False).head(20)
    months = sorted([m for m in d["_month"].dropna().unique() if m <= today.strftime("%Y-%m")])[-24:]
    mcount = d[d["_month"].isin(months)].groupby("_month").size().reindex(months, fill_value=0)
    msol = sol[sol["_month"].isin(months)].groupby("_month").size().reindex(months, fill_value=0)
    recent = d[d["_month"].isin(months[-6:])]
    fcount = recent.groupby("분야").size().sort_values(ascending=False)
    comp_rows = d[d["경쟁사"] == "Y"].copy()
    comp_rows["경쟁사명"] = comp_rows["경쟁사명"].replace("", "기타")
    ccount = comp_rows.groupby("경쟁사명").size().sort_values(ascending=False)

    # ---- 시트 ----
    dash = wb.create_sheet("대시보드")
    for col, w in zip("ABCDEFGHIJKLMNOPQRS", [28, 10, 10, 3] + [9] * 15):
        dash.column_dimensions[col].width = w
    dash.append([c(dash, "공공·대학 IT 수주 현황 대시보드", big)])
    dash.append([f"기준 {today.strftime('%Y-%m-%d')} · 전체 {len(d):,}건 · 숫자를 누르면 관련 사업 목록으로 이동합니다."])
    dash.append([])
    # 표1: 솔루션 연관 업체 TOP 20
    r0 = 4
    dash.append([c(dash, "① 자사 솔루션 연관 사업을 많이 수주한 업체 TOP 20 (에스티씨랩 제외)", bold)])
    dash.append([c(dash, "업체", hf, navy), c(dash, "건수", hf, navy), c(dash, "진행중", hf, navy)])
    t1_first = r0 + 2
    drill_comp_rows = sol.copy()
    gk_to_name = top_comp["업체명"].to_dict()
    drill_comp_rows["_drill"] = drill_comp_rows["_gk"].map(gk_to_name)
    a1 = grouped_table("상세_연관업체", drill_comp_rows[drill_comp_rows["_drill"].notna()], "_drill", list(top_comp["업체명"]),
                       "대시보드 ① 업체별 사업 목록 (자사 솔루션·API·경쟁사·예약/대기열 관련)")
    for _, r in top_comp.iterrows():
        nm = r["업체명"]
        dash.append([nm, c(dash, f'=HYPERLINK("#\'상세_연관업체\'!A{a1.get(nm, 1)}",{int(r["건수"])})', font=link_font), int(r["진행"])])
    t1_last = t1_first + len(top_comp) - 1
    ch = BarChart()
    ch.type, ch.style, ch.title = "bar", 10, "자사 솔루션 연관 수주 TOP 20 (건)"
    ch.add_data(Reference(dash, min_col=2, min_row=t1_first - 1, max_row=t1_last), titles_from_data=True)
    ch.set_categories(Reference(dash, min_col=1, min_row=t1_first, max_row=t1_last))
    ch.y_axis.majorGridlines = None
    ch.x_axis.scaling.orientation = "maxMin"
    ch.dataLabels = _labels()
    ch.legend = None
    ch.series[0].graphicalProperties.solidFill = BAR
    ch.height, ch.width = 11, 18
    dash.add_chart(ch, f"E{r0}")

    # 표2: 월별 신규 사업
    while True:
        cur = r0 + 2 + len(top_comp)
        break
    pad = max(0, 26 - len(top_comp))
    for _ in range(pad):
        dash.append([])
    r2 = cur + pad + 1
    dash.append([])
    dash.append([c(dash, "② 월별 신규 사업 (최근 24개월, 계약·낙찰·공고일 기준)", bold)])
    dash.append([c(dash, "월", hf, navy), c(dash, "전체", hf, navy), c(dash, "솔루션 연관", hf, navy)])
    t2_first = r2 + 2
    a2 = grouped_table("상세_월별", d[d["_month"].isin(months)], "_month", list(reversed(months)), "대시보드 ② 월별 신규 사업 (최근 월부터)")
    for m in months:
        dash.append([m, c(dash, f'=HYPERLINK("#\'상세_월별\'!A{a2.get(m, 1)}",{int(mcount[m])})', font=link_font), int(msol[m])])
    t2_last = t2_first + len(months) - 1
    ch2 = BarChart()
    ch2.type, ch2.style, ch2.title = "col", 10, "월별 신규 IT 사업 (건)"
    ch2.add_data(Reference(dash, min_col=2, min_row=t2_first - 1, max_row=t2_last), titles_from_data=True)
    ch2.set_categories(Reference(dash, min_col=1, min_row=t2_first, max_row=t2_last))
    ch2.dataLabels = _labels()
    ch2.legend = None
    ch2.y_axis.majorGridlines = None
    ch2.series[0].graphicalProperties.solidFill = BAR
    ch2.height, ch2.width = 10, 26
    dash.add_chart(ch2, f"E{r2 + 1}")
    ch2b = BarChart()
    ch2b.type, ch2b.style, ch2b.title = "col", 10, "월별 솔루션 연관 신규 사업 (건)"
    ch2b.add_data(Reference(dash, min_col=3, min_row=t2_first - 1, max_row=t2_last), titles_from_data=True)
    ch2b.set_categories(Reference(dash, min_col=1, min_row=t2_first, max_row=t2_last))
    ch2b.dataLabels = _labels()
    ch2b.legend = None
    ch2b.y_axis.majorGridlines = None
    ch2b.series[0].graphicalProperties.solidFill = "1B8A6B"
    ch2b.height, ch2b.width = 10, 26
    dash.add_chart(ch2b, f"E{r2 + 22}")
    for _ in range(max(0, 42 - len(months))):
        dash.append([])

    # 표3: 분야별 (최근 6개월)
    r3 = t2_last + max(0, 42 - len(months)) + 2
    dash.append([])
    dash.append([c(dash, f"③ 최근 6개월 신규 사업 분야 ({months[-6] if len(months) >= 6 else ''} ~ {months[-1] if months else ''})", bold)])
    dash.append([c(dash, "분야", hf, navy), c(dash, "건수", hf, navy)])
    t3_first = r3 + 2
    a3 = grouped_table("상세_분야", recent, "분야", list(fcount.index), "대시보드 ③ 최근 6개월 분야별 신규 사업")
    for f_, n_ in fcount.items():
        dash.append([f_, c(dash, f'=HYPERLINK("#\'상세_분야\'!A{a3.get(f_, 1)}",{int(n_)})', font=link_font)])
    t3_last = t3_first + len(fcount) - 1
    ch3 = BarChart()
    ch3.type, ch3.style, ch3.title = "bar", 10, "최근 6개월 분야별 신규 사업 (건)"
    ch3.add_data(Reference(dash, min_col=2, min_row=t3_first - 1, max_row=t3_last), titles_from_data=True)
    ch3.set_categories(Reference(dash, min_col=1, min_row=t3_first, max_row=t3_last))
    ch3.x_axis.scaling.orientation = "maxMin"
    ch3.dataLabels = _labels()
    ch3.legend = None
    ch3.y_axis.majorGridlines = None
    ch3.series[0].graphicalProperties.solidFill = BAR
    ch3.height, ch3.width = 9, 18
    dash.add_chart(ch3, f"E{r3 + 1}")
    for _ in range(max(0, 20 - len(fcount))):
        dash.append([])

    # 표4: 경쟁사별
    r4 = t3_last + max(0, 20 - len(fcount)) + 2
    dash.append([])
    dash.append([c(dash, "④ 경쟁사별 수주·납품 (전체 기간)", bold)])
    dash.append([c(dash, "경쟁사", hf, navy), c(dash, "건수", hf, navy)])
    t4_first = r4 + 2
    a4 = grouped_table("경쟁사 수주", comp_rows, "경쟁사명", list(ccount.index),
                       "경쟁사(데브와이·스크립터스·에버스핀·소프트베이스·가온아이)가 수주했거나 그 제품(DynaPath·에버세이프·xQueue)이 들어간 사업")
    for k_, n_ in ccount.items():
        dash.append([k_, c(dash, f'=HYPERLINK("#\'경쟁사 수주\'!A{a4.get(k_, 1)}",{int(n_)})', font=link_font)])
    if len(ccount):
        t4_last = t4_first + len(ccount) - 1
        ch4 = BarChart()
        ch4.type, ch4.style, ch4.title = "bar", 10, "경쟁사별 건수"
        ch4.add_data(Reference(dash, min_col=2, min_row=t4_first - 1, max_row=t4_last), titles_from_data=True)
        ch4.set_categories(Reference(dash, min_col=1, min_row=t4_first, max_row=t4_last))
        ch4.x_axis.scaling.orientation = "maxMin"
        ch4.dataLabels = _labels()
        ch4.legend = None
        ch4.y_axis.majorGridlines = None
        ch4.series[0].graphicalProperties.solidFill = "C0392B"
        ch4.height, ch4.width = 7, 18
        dash.add_chart(ch4, f"E{r4 + 1}")

    # 업체검색
    ws = wb.create_sheet("업체검색")
    for i, (_, _, w) in enumerate(VIEW, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    T = "'통합 내역'"
    hcol = get_column_letter(len(VIEW) + len(HELP) + 1)
    ws.append([c(ws, "업체 검색", big)])
    ws.append(["노란 칸(B3)에 업체명 일부 입력 → 그 업체 사업 전부 (종료 임박·진행중 먼저). (주)·띄어쓰기·대소문자·한글/영문 표기 차이 무시. 사업명을 누르면 원문."])
    ws.append([c(ws, "업체명 →", bold), c(ws, "", fill=yellow), c(ws, "결과 건수", bold),
               f"=IF(F3=\"\",\"\",COUNT({T}!{hcol}:{hcol}))", c(ws, "인식한 검색어", bold), f'={ce.excel_norm_formula("B3")}'])
    ws.append([c(ws, t, hf, navy) for t, _, _ in VIEW])
    ws.freeze_panes = "F5"
    for k in range(1, 1001):
        pos = f"SMALL({T}!${hcol}:${hcol},{k})"
        row = []
        for t, key, _ in VIEW:
            col = L[t]
            if key == "사업명":
                u = L["원문 주소"]
                row.append(f'=IFERROR(IF(INDEX({T}!{u}:{u},{pos})="",INDEX({T}!{col}:{col},{pos})&"",'
                           f'HYPERLINK(INDEX({T}!{u}:{u},{pos}),INDEX({T}!{col}:{col},{pos})&"")),"")')
            elif key == "금액":
                w = L["금액(원)"]
                row.append(f'=IFERROR(IF(INDEX({T}!{w}:{w},{pos})="","",IF(INDEX({T}!{w}:{w},{pos})>=10000000,'
                           f'TEXT(INDEX({T}!{w}:{w},{pos})/100000000,"#,##0.0")&"억",TEXT(INDEX({T}!{w}:{w},{pos}),"#,##0"))),"")')
            else:
                row.append(f'=IFERROR(INDEX({T}!{col}:{col},{pos})&"","")')
        ws.append(row)

    # 업체별 요약 (같은 앞 5칸 기준)
    g = d[d["대표업체"] != ""].groupby("_gk")
    summ = pd.DataFrame({
        "업체명": g["대표업체"].agg(lambda x: x.value_counts().index[0]),
        "주 지역": g["지역"].agg(lambda x: x.value_counts().index[0]),
        "공공": g["기관특성"].agg(lambda x: int((x == "공공").sum())),
        "대학(국·공립)": g["대학구분"].agg(lambda x: int((x == "국·공립").sum())),
        "대학(사립)": g["대학구분"].agg(lambda x: int((x == "사립").sum())),
        "진행중": g["종료"].agg(lambda x: int((x == "진행중").sum())),
        "종료 임박": g["종료"].agg(lambda x: int((x == "종료 임박").sum())),
        "종료": g["종료"].agg(lambda x: int((x == "종료").sum())),
        "솔루션 연관": g["솔루션연관"].agg(lambda x: int((x == "Y").sum())),
        "금액 합계": g["금액"].sum(min_count=1),
        "가장 가까운 종료": g.apply(lambda x: x.loc[x["_days"] >= 0, "사업종료일"].min() if (x["_days"] >= 0).any() else ""),
        "주요 수요기관": g["수요기관"].agg(lambda x: ", ".join(x.value_counts().index[:3])),
    }).sort_values(["진행중", "종료 임박"], ascending=False)
    ws = wb.create_sheet("업체별 요약")
    cols = list(summ.columns)
    for i, k in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = {"업체명": 28, "주요 수요기관": 46, "가장 가까운 종료": 12, "금액 합계": 11}.get(k, 9)
    ws.append([c(ws, k, hf, navy) for k in cols])
    ws.freeze_panes = "B2"
    for r in summ.to_dict("records"):
        ws.append([money(ws, r[k]) if k == "금액 합계" else r[k] for k in cols])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(summ) + 1}"

    table("솔루션·API 관련", sol, note="자사 솔루션(대기열·예약·매크로·API 트래픽 등)·API·경쟁사 관련 사업")
    univ = d[d["기관특성"] == "대학"]
    table("대학 사업", univ, note=f"대학 사업 {len(univ):,}건 (국·공립 {int((univ['대학구분'] == '국·공립').sum()):,} · 사립 {int((univ['대학구분'] == '사립').sum()):,})")
    table("통합 내역", d, helper=True)

    # 안내 + 대학 수집 결과
    ws = wb.create_sheet("안내")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 110
    n = d["출처"].value_counts()
    for a, b in [
        ("공공·대학 IT 수주 통합 현황", ""), ("기준일", today.strftime("%Y-%m-%d")),
        ("출처", f"나라장터 계약 {int(n.get('계약', 0)):,} · 낙찰 {int(n.get('낙찰', 0)):,} · 쇼핑몰 납품 {int(n.get('쇼핑몰 납품', 0)):,} · 대학 홈페이지·전자입찰 {int(n.get('대학 홈페이지', 0)):,}"),
        ("열 구성", "모든 표 공통: 기관특성 · 대학구분 · 업체명 · 수요기관 · 지역(고정) | 사업명 · 제품 · 금액 · 사업종료일 · 종료 · 종료 D-day"),
        ("종료", f"진행중 / 종료 임박(종료일까지 {ce.SOON_DAYS}일 이내, 빨간 칸) / 종료 / 미상"),
        ("금액", "천만원 이상 '억'(예: 10.0억, 0.5억), 미만은 원 단위(예: 9,500,000)"),
        ("사업종료일", "계약은 계약 종료일. 낙찰·쇼핑몰 납품·대학 공고는 종료일 정보가 없어 일자 + 1년으로 추정"),
        ("경쟁사", "업체: 데브와이·스크립터스·에버스핀·소프트베이스·가온아이 / 제품: DynaPath(스크립터스)·에버세이프(에버스핀)·xQueue(소프트베이스)"),
        ("지역", "수요기관 소재 시·도. 확인 못 한 기관은 '미확인'"),
    ]:
        ws.append([c(ws, a, bold), b])
    rep, _ = store.load_cache("univ_crawl_report")
    if rep:
        ws = wb.create_sheet("대학 수집 결과")
        for i, w in enumerate([24, 12, 8, 8, 60, 40], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.append([c(ws, k, hf, navy) for k in ("학교", "결과", "3년 내 글", "IT 글", "입찰 게시판 주소", "홈페이지")])
        for r in sorted(rep, key=lambda x: (x.get("status") != "성공", -int(x.get("it_posts") or 0))):
            ws.append([clean(r.get("school")), r.get("status"), r.get("posts"), r.get("it_posts"), clean(r.get("board")), clean(r.get("home"))])

    buf = io.BytesIO()
    wb.save(buf)
    meta = {"rows": len(d), "by_source": {k: int(v) for k, v in n.items()}, "companies": int(len(summ)),
            "univ_rows": int(len(univ)), "univ_private": int((univ["대학구분"] == "사립").sum()),
            "competitor_rows": int(len(comp_rows)), "solution_rows": int(len(sol))}
    return buf.getvalue(), meta


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="integrated.xlsx")
    ap.add_argument("--dlvr-cache", default="raw_cache.json")
    a = ap.parse_args()
    comp = store.load_competitors()[0]
    load_code_region()
    cons, ntce = load_contracts(comp)
    awards = load_awards(comp, ntce)
    dlvr = load_dlvr(comp, a.dlvr_cache)
    univ = load_univ_bids(comp)
    parts = [x for x in (cons, awards, dlvr, univ) if not x.empty]
    if not parts:
        print("[SKIP] 자료 없음")
        return 1
    all_df = pd.concat(parts, ignore_index=True)
    data, meta = build(all_df)
    open(a.out, "wb").write(data)
    import base64
    fname = f"공공IT_수주통합현황_{datetime.now().strftime('%Y%m%d')}.xlsx"
    try:
        store.save_cache("integrated_excel", {"filename": fname, "b64": base64.b64encode(data).decode(), **meta})
        store.save_cache("integrated_excel_meta", {"filename": fname, "date": datetime.now().strftime("%Y-%m-%d"), "size": len(data), **meta})
    except Exception as e:
        print(f"[WARN] 대시보드용 저장 실패: {type(e).__name__}")
    print(f"[OK] 통합 엑셀: {len(data) / 1e6:.1f}MB · {json.dumps(meta, ensure_ascii=False)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
